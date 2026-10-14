"""Mutfak zayi testleri:  python manage.py test panel.tests_zayi"""
import datetime
import io
from decimal import Decimal

from django.apps import apps as global_apps
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from . import zayi_rapor
from .models import (MutfakZayi, MutfakZayiKalemi, Personel, Rol, Sube, Vardiya, VardiyaTipi,
                     zayi_urun_anahtar, zayi_urun_baslik)
from .zayi_kapsam import zayi_menusu_gorunur, zayi_subeleri
from .zayi_views import miktar_oku

_AYAR = dict(
    SECURE_SSL_REDIRECT=False,
    STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
              'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}},
)


class MiktarVeAdTestleri(TestCase):
    def test_miktar_oku(self):
        self.assertEqual(miktar_oku('100'), Decimal('100.00'))
        self.assertEqual(miktar_oku('1,5'), Decimal('1.50'))
        self.assertEqual(miktar_oku('1.5'), Decimal('1.50'))
        self.assertEqual(miktar_oku('1.250'), Decimal('1250.00'))        # Türkçe binlik
        self.assertEqual(miktar_oku('0.125'), Decimal('0.13'))           # ondalık (0 ile başlıyor)
        self.assertEqual(miktar_oku('1.250,5'), Decimal('1250.50'))
        for kotu in ('', '   ', '0', '-3', 'abc', '1e9', 'NaN', 'inf', '2000000'):
            self.assertIsNone(miktar_oku(kotu), kotu)

    def test_urun_anahtar_turkce(self):
        self.assertEqual(zayi_urun_anahtar('  Domates  '), 'domates')
        self.assertEqual(zayi_urun_anahtar('DOMATES'), 'domates')
        self.assertEqual(zayi_urun_anahtar('İÇ PİLAV'), 'iç pilav')
        self.assertEqual(zayi_urun_anahtar('Pesto   Sos'), 'pesto sos')
        self.assertEqual(zayi_urun_baslik('pesto sos'), 'Pesto Sos')
        self.assertEqual(zayi_urun_baslik('ışık ipi'), 'Işık İpi')


class ZayiBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.a = Sube.objects.create(ad='A Şube', mutfak_var=True)
        cls.b = Sube.objects.create(ad='B Şube', mutfak_var=True)
        cls.c = Sube.objects.create(ad='C Şube (mutfaksız)', mutfak_var=False)
        cls.depo = Sube.objects.create(ad='Depo', depo_mu=True, mutfak_var=True)
        cls.bugun = timezone.localdate()

    def kisi(self, rol, ad=None, sube=None, sorumlu=()):
        u = User.objects.create(username='u%d' % (User.objects.count() + 1))
        p = Personel.objects.create(user=u, ad_soyad=ad or rol, rol=rol, sube=sube)
        for s in sorumlu:
            p.sorumlu_subeler.add(s)
        return p

    def gir(self, p):
        self.client.force_login(p.user)

    def kalem(self, sube, urun, miktar, birim, tarih=None, personel=None, aciklama=''):
        return MutfakZayiKalemi(sube=sube, personel=personel, personel_ad_arsiv=personel.ad_soyad if personel else 'x',
                                tarih=tarih or self.bugun, urun=urun, miktar=Decimal(str(miktar)),
                                birim=birim, aciklama=aciklama)

    def kaydet(self, *a, **k):
        x = self.kalem(*a, **k)
        x.save()
        return x


@override_settings(**_AYAR)
class KapsamTestleri(ZayiBase):
    def test_roller_ve_kapsam(self):
        gm = self.kisi(Rol.GENEL_MUDUR)
        op = self.kisi(Rol.OPERATOR)
        yat = self.kisi(Rol.YATIRIMCI)
        sor = self.kisi(Rol.MUTFAK_SORUMLUSU)
        for p in (gm, op, yat, sor):
            self.assertEqual({s.id for s in zayi_subeleri(p)}, {self.a.id, self.b.id}, p.rol)   # depo ve mutfaksız yok
        mud = self.kisi(Rol.MUDUR, sorumlu=[self.b, self.c])
        self.assertEqual([s.id for s in zayi_subeleri(mud)], [self.b.id])
        mm = self.kisi(Rol.MAGAZA_MUDURU, sube=self.a)
        self.assertEqual([s.id for s in zayi_subeleri(mm)], [self.a.id])
        mm2 = self.kisi(Rol.MAGAZA_MUDURU, sube=self.a, sorumlu=[self.a, self.b])
        self.assertEqual({s.id for s in zayi_subeleri(mm2)}, {self.a.id, self.b.id})
        sef = self.kisi(Rol.SEF, sube=self.b)
        self.assertEqual([s.id for s in zayi_subeleri(sef)], [self.b.id])
        sef_mutfaksiz = self.kisi(Rol.SEF, sube=self.c)
        self.assertEqual(zayi_subeleri(sef_mutfaksiz), [])
        self.assertFalse(zayi_menusu_gorunur(sef_mutfaksiz))
        self.assertTrue(zayi_menusu_gorunur(sef))
        for rol in (Rol.PERSONEL, Rol.EGITMEN, Rol.SATIN_ALMA, Rol.SEVKIYAT):
            self.assertEqual(zayi_subeleri(self.kisi(rol, sube=self.a)), [])
            self.assertFalse(zayi_menusu_gorunur(self.kisi(rol, sube=self.a)))

    def test_mutfak_personeli_atama_subesi(self):
        mp = self.kisi(Rol.MUTFAK_PERSONEL, sube=self.c)                       # kendi şubesi mutfaksız
        self.assertEqual(zayi_subeleri(mp), [])
        Vardiya.objects.create(personel=mp, tarih=self.bugun, vardiya_tipi=VardiyaTipi.MUTFAK_GOREVI,
                               atanan_sube=self.b)
        self.assertEqual([s.id for s in zayi_subeleri(mp)], [self.b.id])        # bugünkü atama kazanır

    def test_menude_baglanti(self):
        sef = self.kisi(Rol.SEF, sube=self.a)
        sef_yok = self.kisi(Rol.SEF, sube=self.c)
        gm = self.kisi(Rol.GENEL_MUDUR)
        href = 'href="%s"' % reverse('mutfak_zayi')
        self.gir(sef)
        self.assertContains(self.client.get(reverse('ana_sayfa')), href)
        self.gir(sef_yok)
        self.assertNotContains(self.client.get(reverse('ana_sayfa')), href)
        self.gir(gm)
        self.assertContains(self.client.get(reverse('gosterge')), href)
        # tüm şubelerde mutfak kapanırsa genel müdür de görmez
        Sube.objects.update(mutfak_var=False)
        self.assertNotContains(self.client.get(reverse('gosterge')), href)


@override_settings(**_AYAR)
class GirisTestleri(ZayiBase):
    def post_ekle(self, satirlar, **ek):
        veri = {'islem': 'zayi_ekle', 'urun': [], 'miktar': [], 'birim': [], 'aciklama': []}
        for u, m, b, a in satirlar:
            veri['urun'].append(u); veri['miktar'].append(m); veri['birim'].append(b); veri['aciklama'].append(a)
        veri.update(ek)
        return self.client.post(reverse('mutfak_zayi'), veri)

    def test_personel_yazar(self):
        mp = self.kisi(Rol.MUTFAK_PERSONEL, 'Ali', sube=self.a)
        self.gir(mp)
        self.assertEqual(self.client.get(reverse('mutfak_zayi')).status_code, 200)
        r = self.post_ekle([('Domates', '100', 'gr', ''), ('Pesto Sos', '500', 'ml', 'kapağı açık kaldı'),
                            ('', '', 'gr', '')])                                # boş satır atlanır
        self.assertEqual(r.status_code, 302)
        k = list(MutfakZayiKalemi.objects.order_by('id'))
        self.assertEqual(len(k), 2)
        self.assertEqual((k[0].sube_id, k[0].urun, k[0].miktar, k[0].birim, k[0].tarih),
                         (self.a.id, 'Domates', Decimal('100.00'), 'gr', self.bugun))
        self.assertEqual((k[0].olcu, k[0].taban_miktar, k[0].personel_ad_arsiv), ('agirlik', Decimal('100.00'), 'Ali'))
        self.assertEqual((k[1].olcu, k[1].aciklama), ('hacim', 'kapağı açık kaldı'))
        # liste ekranında görünür
        sayfa = self.client.get(reverse('mutfak_zayi'))
        self.assertContains(sayfa, 'Domates')
        self.assertContains(sayfa, 'kapağı açık kaldı')

    def test_personel_sube_ve_tarih_secemez(self):
        mp = self.kisi(Rol.MUTFAK_PERSONEL, sube=self.a)
        self.gir(mp)
        dun = (self.bugun - datetime.timedelta(days=3)).strftime('%Y-%m-%d')
        self.post_ekle([('Soğan', '1', 'kg', '')], sube_id=str(self.b.id), tarih=dun)
        k = MutfakZayiKalemi.objects.get()
        self.assertEqual((k.sube_id, k.tarih), (self.a.id, self.bugun))        # POST ile oynanamaz

    def test_mutfaksiz_sube_engellenir(self):
        mp = self.kisi(Rol.MUTFAK_PERSONEL, sube=self.c)
        self.gir(mp)
        sayfa = self.client.get(reverse('mutfak_zayi'))
        self.assertContains(sayfa, 'mutfağı olan bir şubeye atamanız görünmüyor')
        self.assertNotContains(sayfa, 'name="miktar"')
        self.post_ekle([('Domates', '100', 'gr', '')])
        self.assertEqual(MutfakZayiKalemi.objects.count(), 0)

    def test_hatali_satir_hicbir_sey_kaydetmez(self):
        mp = self.kisi(Rol.MUTFAK_PERSONEL, sube=self.a)
        self.gir(mp)
        r = self.post_ekle([('Domates', '100', 'gr', ''), ('Biber', 'abc', 'gr', ''), ('Süt', '2', 'galon', '')])
        self.assertEqual(r.status_code, 200)                                    # form geri gelir
        self.assertEqual(MutfakZayiKalemi.objects.count(), 0)
        self.assertContains(r, 'miktar geçersiz')
        self.assertContains(r, 'birim seçin')
        self.assertContains(r, 'value="Domates"')                               # yazılanlar korunur

    def test_ust_sinir(self):
        mp = self.kisi(Rol.MUTFAK_PERSONEL, sube=self.a)
        self.gir(mp)
        self.post_ekle([('Ürün %d' % i, '1', 'adet', '') for i in range(41)])
        self.assertEqual(MutfakZayiKalemi.objects.count(), 0)
        self.post_ekle([('Ürün %d' % i, '1', 'adet', '') for i in range(40)])
        self.assertEqual(MutfakZayiKalemi.objects.count(), 40)

    def test_sorumlu_sube_ve_tarih_secer(self):
        sor = self.kisi(Rol.MUTFAK_SORUMLUSU, 'Sorumlu')
        self.gir(sor)
        dun = (self.bugun - datetime.timedelta(days=1)).strftime('%Y-%m-%d')
        self.post_ekle([('Tavuk', '2,5', 'kg', '')], sube_id=str(self.b.id), tarih=dun)
        k = MutfakZayiKalemi.objects.get()
        self.assertEqual((k.sube_id, k.tarih, k.miktar), (self.b.id, self.bugun - datetime.timedelta(days=1), Decimal('2.50')))
        # mutfaksız / depo şubesine yazamaz, şube seçmeden de yazamaz
        self.post_ekle([('Tavuk', '1', 'kg', '')], sube_id=str(self.c.id), tarih=dun)
        self.post_ekle([('Tavuk', '1', 'kg', '')], sube_id=str(self.depo.id), tarih=dun)
        self.post_ekle([('Tavuk', '1', 'kg', '')], tarih=dun)
        # gelecek tarih ve çok eski tarih olmaz
        yarin = (self.bugun + datetime.timedelta(days=1)).strftime('%Y-%m-%d')
        eski = (self.bugun - datetime.timedelta(days=61)).strftime('%Y-%m-%d')
        self.post_ekle([('Tavuk', '1', 'kg', '')], sube_id=str(self.a.id), tarih=yarin)
        self.post_ekle([('Tavuk', '1', 'kg', '')], sube_id=str(self.a.id), tarih=eski)
        self.assertEqual(MutfakZayiKalemi.objects.count(), 1)

    def test_yazma_yetkisi_olmayanlar(self):
        for rol, kw in ((Rol.SEF, dict(sube=self.a)), (Rol.MAGAZA_MUDURU, dict(sube=self.a)),
                        (Rol.GENEL_MUDUR, {}), (Rol.OPERATOR, {}), (Rol.YATIRIMCI, {})):
            self.gir(self.kisi(rol, **kw))
            sayfa = self.client.get(reverse('mutfak_zayi'))
            self.assertEqual(sayfa.status_code, 200, rol)
            self.assertNotContains(sayfa, 'name="miktar"', msg_prefix=rol)
            self.post_ekle([('Domates', '100', 'gr', '')], sube_id=str(self.a.id))
        self.assertEqual(MutfakZayiKalemi.objects.count(), 0)

    def test_yetkisiz_roller_sayfayi_goremez(self):
        for rol in (Rol.PERSONEL, Rol.EGITMEN, Rol.SATIN_ALMA, Rol.SEVKIYAT):
            self.gir(self.kisi(rol, sube=self.a))
            self.assertEqual(self.client.get(reverse('mutfak_zayi')).status_code, 302, rol)
            self.assertEqual(self.client.get(reverse('mutfak_zayi_rapor')).status_code, 302, rol)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('mutfak_zayi')).status_code, 302)

    def test_kayit_gunluk_gorunum_kapsami(self):
        self.kaydet(self.a, 'Domates', 1, 'kg')
        self.kaydet(self.b, 'Biber', 2, 'kg')
        self.kaydet(self.c, 'Gizli', 3, 'kg')                                    # mutfaksız şubenin eski kaydı
        sef = self.kisi(Rol.SEF, sube=self.a)
        self.gir(sef)
        s = self.client.get(reverse('mutfak_zayi'))
        self.assertContains(s, 'Domates'); self.assertNotContains(s, 'Biber'); self.assertNotContains(s, 'Gizli')
        # başka şubeyi parametreyle de göremez
        s = self.client.get(reverse('mutfak_zayi') + '?sube=%d' % self.b.id)
        self.assertNotContains(s, 'Biber'); self.assertContains(s, 'Domates')
        self.gir(self.kisi(Rol.MUTFAK_SORUMLUSU))
        s = self.client.get(reverse('mutfak_zayi'))
        self.assertContains(s, 'Domates'); self.assertContains(s, 'Biber'); self.assertNotContains(s, 'Gizli')

    def test_eski_fotografli_kayit_gorunur(self):
        MutfakZayi.objects.create(sube=self.a, personel_ad_arsiv='Eski', foto='mutfak_zayi/x.jpg', aciklama='eski not')
        self.gir(self.kisi(Rol.SEF, sube=self.a))
        self.assertContains(self.client.get(reverse('mutfak_zayi')), 'Eski fotoğraflı kayıtlar')


@override_settings(**_AYAR)
class SilmeTestleri(ZayiBase):
    def sil(self, k):
        return self.client.post(reverse('mutfak_zayi'), {'islem': 'zayi_sil', 'kalem_id': str(k.id)})

    def test_personel_yalniz_kendi_bugunku_kaydini_siler(self):
        ali = self.kisi(Rol.MUTFAK_PERSONEL, 'Ali', sube=self.a)
        veli = self.kisi(Rol.MUTFAK_PERSONEL, 'Veli', sube=self.a)
        benim = self.kaydet(self.a, 'Domates', 1, 'kg', personel=ali)
        onun = self.kaydet(self.a, 'Biber', 1, 'kg', personel=veli)
        eski = self.kaydet(self.a, 'Süt', 1, 'lt', personel=ali, tarih=self.bugun - datetime.timedelta(days=1))
        self.gir(ali)
        self.sil(onun); self.sil(eski)
        self.assertTrue(MutfakZayiKalemi.objects.filter(id=onun.id).exists())
        self.assertTrue(MutfakZayiKalemi.objects.filter(id=eski.id).exists())
        self.sil(benim)
        self.assertFalse(MutfakZayiKalemi.objects.filter(id=benim.id).exists())

    def test_sorumlu_siler_digerleri_silemez(self):
        ali = self.kisi(Rol.MUTFAK_PERSONEL, sube=self.a)
        k1 = self.kaydet(self.a, 'Domates', 1, 'kg', personel=ali)
        k2 = self.kaydet(self.c, 'Gizli', 1, 'kg', personel=ali)                 # kapsam dışı (mutfaksız)
        for rol, kw in ((Rol.SEF, dict(sube=self.a)), (Rol.MAGAZA_MUDURU, dict(sube=self.a)),
                        (Rol.GENEL_MUDUR, {}), (Rol.OPERATOR, {}), (Rol.YATIRIMCI, {})):
            self.gir(self.kisi(rol, **kw))
            self.sil(k1)
        self.assertTrue(MutfakZayiKalemi.objects.filter(id=k1.id).exists())
        self.gir(self.kisi(Rol.MUTFAK_SORUMLUSU))
        self.sil(k2)
        self.assertTrue(MutfakZayiKalemi.objects.filter(id=k2.id).exists())
        self.sil(k1)
        self.assertFalse(MutfakZayiKalemi.objects.filter(id=k1.id).exists())


@override_settings(**_AYAR)
class RaporTestleri(ZayiBase):
    def setUp(self):
        t = datetime.date(2026, 10, 14)                                          # Çarşamba
        self.t = t
        k = self.kaydet
        k(self.a, 'Domates', 100, 'gr', tarih=t)
        k(self.a, 'domates ', 2, 'kg', tarih=t)                                  # aynı ürün, farklı yazım/birim
        k(self.a, 'Süt', 500, 'ml', tarih=t, aciklama='döküldü')
        k(self.a, 'Süt', 1, 'lt', tarih=t + datetime.timedelta(days=1))
        k(self.a, 'Kase', 3, 'adet', tarih=t + datetime.timedelta(days=2))
        k(self.b, 'Domates', 400, 'gr', tarih=t)
        k(self.b, 'Domates', 50, 'gr', tarih=datetime.date(2026, 10, 20))        # sonraki hafta
        k(self.b, 'Tavuk', 1, 'kg', tarih=datetime.date(2026, 9, 30))            # önceki ay
        k(self.c, 'Gizli', 9, 'kg', tarih=t)                                     # mutfaksız şube

    def veri(self, subeler, b, e):
        return zayi_rapor.rapor_verisi(subeler, b, e)

    def test_donemler(self):
        t = self.t
        d = zayi_rapor.donem_hesapla
        self.assertEqual(d('gun', t)[1:3], (t, t))
        h = d('hafta', t)
        self.assertEqual((h[1], h[2]), (datetime.date(2026, 10, 12), datetime.date(2026, 10, 18)))
        a = d('ay', t)
        self.assertEqual((a[1], a[2], a[3]), (datetime.date(2026, 10, 1), datetime.date(2026, 10, 31), 'Ekim 2026'))
        self.assertEqual(d('ay', datetime.date(2026, 12, 5))[2], datetime.date(2026, 12, 31))
        self.assertEqual(d('ay', datetime.date(2026, 2, 5))[2], datetime.date(2026, 2, 28))
        # özel aralık: ters verilirse düzelir, 366 günü aşarsa kısaltılır + uyarı
        o = d('ozel', t, datetime.date(2026, 10, 20), datetime.date(2026, 10, 10))
        self.assertEqual((o[1], o[2], o[6]), (datetime.date(2026, 10, 10), datetime.date(2026, 10, 20), ''))
        u = d('ozel', t, datetime.date(2020, 1, 1), datetime.date(2026, 1, 1))
        self.assertEqual((u[2] - u[1]).days + 1, 366)
        self.assertTrue(u[6])
        # onceki / sonraki
        self.assertEqual(d('hafta', t)[4], datetime.date(2026, 10, 5))
        self.assertEqual(d('ay', t)[4], datetime.date(2026, 9, 1))
        self.assertEqual(d('ay', t)[5], datetime.date(2026, 11, 1))
        self.assertEqual(d('gecersiz', t)[0], 'gun')

    def test_toplamlar(self):
        v = self.veri([self.a, self.b], datetime.date(2026, 10, 12), datetime.date(2026, 10, 18))
        oz = {o['sube'].ad: o for o in v['ozet']}
        a, b = oz['A Şube'], oz['B Şube']
        self.assertEqual(a['top']['agirlik'], Decimal('2100'))                    # 100 gr + 2 kg
        self.assertEqual(a['top']['hacim'], Decimal('1500'))                      # 500 ml + 1 lt
        self.assertEqual(a['top']['adet'], Decimal('3'))
        self.assertEqual(a['kayit'], 5)
        self.assertEqual((a['yazi']['agirlik'], a['yazi']['hacim'], a['yazi']['adet']), ('2,1 kg', '1,5 lt', '3 adet'))
        self.assertEqual(b['top']['agirlik'], Decimal('400'))                     # 20 Ekim ve eylül dışarıda
        self.assertEqual(b['yazi']['agirlik'], '400 gr')
        self.assertEqual(v['genel']['kayit'], 6)
        self.assertEqual(v['genel']['top']['agirlik'], Decimal('2500'))
        # ürünler: Domates tek satır (büyük/küçük harf ve kg/gr birleşti), Süt ayrı ölçüde
        urunler = {(u['ad'], u['olcu']): u for u in v['urunler']}
        dom = urunler[('Domates', 'agirlik')]
        self.assertEqual(dom['toplam'], Decimal('2500'))
        self.assertEqual(dom['subeler'], {self.a.id: Decimal('2100'), self.b.id: Decimal('400')})
        self.assertEqual(dom['toplam_yazi'], '2,5 kg')
        self.assertNotIn(('Gizli', 'agirlik'), urunler)
        self.assertEqual([g['tarih'].day for g in v['gunler']], [14, 15, 16])

    def test_miktar_yaz(self):
        self.assertEqual(zayi_rapor.miktar_yaz(Decimal('999.5'), 'agirlik'), '999,5 gr')
        self.assertEqual(zayi_rapor.miktar_yaz(Decimal('1000'), 'agirlik'), '1 kg')
        self.assertEqual(zayi_rapor.miktar_yaz(Decimal('1250'), 'hacim'), '1,25 lt')
        self.assertEqual(zayi_rapor.miktar_yaz(Decimal('3'), 'adet'), '3 adet')

    def url(self, **q):
        return reverse('mutfak_zayi_rapor') + ('?' + '&'.join('%s=%s' % kv for kv in q.items()) if q else '')

    def test_kapsam_ve_yetki(self):
        sef = self.kisi(Rol.SEF, sube=self.a)
        self.gir(sef)
        r = self.client.get(self.url(donem='hafta', tarih='2026-10-14'))
        self.assertContains(r, 'Domates'); self.assertNotContains(r, 'Gizli')
        self.assertEqual(r.context['veri']['genel']['top']['agirlik'], Decimal('2100'))   # yalnız kendi şubesi
        # başka şubeyi ?sube= ile isteyince yine yalnız kendi şubesi
        r = self.client.get(self.url(donem='hafta', tarih='2026-10-14', sube=self.b.id))
        self.assertEqual(r.context['veri']['genel']['top']['agirlik'], Decimal('2100'))
        mud = self.kisi(Rol.MUDUR, sorumlu=[self.b])
        self.gir(mud)
        r = self.client.get(self.url(donem='ay', tarih='2026-10-14'))
        self.assertEqual(r.context['veri']['genel']['top']['agirlik'], Decimal('450'))
        # mutfaksız şube sorumlusu: rapor yok
        self.gir(self.kisi(Rol.SEF, sube=self.c))
        self.assertEqual(self.client.get(self.url()).status_code, 302)
        # mutfak personeli rapor alamaz
        self.gir(self.kisi(Rol.MUTFAK_PERSONEL, sube=self.a))
        self.assertEqual(self.client.get(self.url()).status_code, 302)
        self.assertEqual(self.client.get(self.url(bicim='xlsx')).status_code, 302)

    def test_herkes_rapor_alir(self):
        for rol, kw in ((Rol.GENEL_MUDUR, {}), (Rol.OPERATOR, {}), (Rol.YATIRIMCI, {}), (Rol.MUTFAK_SORUMLUSU, {}),
                        (Rol.MUDUR, dict(sorumlu=[self.a])), (Rol.MAGAZA_MUDURU, dict(sube=self.a)),
                        (Rol.SEF, dict(sube=self.a))):
            self.gir(self.kisi(rol, **kw))
            for bicim in ('', 'xlsx', 'pdf'):
                r = self.client.get(self.url(donem='hafta', tarih='2026-10-14', **({'bicim': bicim} if bicim else {})))
                self.assertEqual(r.status_code, 200, (rol, bicim))

    def test_excel(self):
        import openpyxl
        self.gir(self.kisi(Rol.GENEL_MUDUR))
        r = self.client.get(self.url(donem='hafta', tarih='2026-10-14', bicim='xlsx'))
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('mutfak-zayi_20261012_20261018.xlsx', r['Content-Disposition'])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Şube Özeti', 'Ürün x Şube', 'Gün Gün', 'Kayıtlar'])

        def tablo(ws):
            return [[c.value for c in row] for row in ws.iter_rows() if any(c.value is not None for c in row)]
        oz = {row[0]: row for row in tablo(wb['Şube Özeti']) if row[0] in ('A Şube', 'B Şube', 'TOPLAM')}
        self.assertEqual(oz['A Şube'][1:], [5, 2.1, 1.5, 3])
        self.assertEqual(oz['B Şube'][1:], [1, 0.4, 0.0, 0.0])
        self.assertEqual(oz['TOPLAM'][1:], [6, 2.5, 1.5, 3])
        kay = [r_ for r_ in tablo(wb['Kayıtlar'])[1:]]
        self.assertEqual(len(kay), 6)
        self.assertIn('döküldü', [x[6] for x in kay])
        urun = tablo(wb['Ürün x Şube'])
        domates = next(x for x in urun if x[0] == 'Domates')
        self.assertEqual(domates[3:], [2100.0, 400.0, 2500.0])

    def test_pdf(self):
        self.gir(self.kisi(Rol.GENEL_MUDUR))
        r = self.client.get(self.url(donem='ay', tarih='2026-10-14', bicim='pdf'))
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertTrue(r.content.startswith(b'%PDF'))
        self.assertIn('mutfak-zayi_20261001_20261031.pdf', r['Content-Disposition'])
        # Bos dönem de patlamaz
        r = self.client.get(self.url(donem='gun', tarih='2020-01-01', bicim='pdf'))
        self.assertEqual(r.status_code, 200)
        r = self.client.get(self.url(donem='gun', tarih='2020-01-01', bicim='xlsx'))
        self.assertEqual(r.status_code, 200)

    def test_gecersiz_parametreler(self):
        self.gir(self.kisi(Rol.GENEL_MUDUR))
        for q in ('donem=xx&tarih=abc', 'sube=zzz', 'donem=ozel&bas=1&bit=2', 'sube=99999'):
            self.assertEqual(self.client.get(reverse('mutfak_zayi_rapor') + '?' + q).status_code, 200, q)

    def test_sayfa_ozeti(self):
        self.gir(self.kisi(Rol.GENEL_MUDUR))
        r = self.client.get(self.url(donem='hafta', tarih='2026-10-14'))
        self.assertContains(r, 'Şubeler arası karşılaştırma')
        self.assertContains(r, '2,1 kg')
        self.assertContains(r, 'Excel'); self.assertContains(r, 'PDF')


@override_settings(**_AYAR)
class SubeSayfasiTestleri(ZayiBase):
    def test_tik_ayari(self):
        gm = self.kisi(Rol.GENEL_MUDUR)
        self.gir(gm)
        self.assertContains(self.client.get(reverse('sube_yeni')), 'Mutfağı var')
        # aç
        self.client.post(reverse('sube_yeni'), {'islem': 'mutfak_ayar', 'sube_id': self.c.id, 'mutfak_var': 'on'})
        self.c.refresh_from_db(); self.assertTrue(self.c.mutfak_var)
        # kapat (checkbox gönderilmez)
        self.client.post(reverse('sube_yeni'), {'islem': 'mutfak_ayar', 'sube_id': self.c.id})
        self.c.refresh_from_db(); self.assertFalse(self.c.mutfak_var)
        # depo olan şubede açılamaz
        self.depo.mutfak_var = False; self.depo.save()
        self.client.post(reverse('sube_yeni'), {'islem': 'mutfak_ayar', 'sube_id': self.depo.id, 'mutfak_var': 'on'})
        self.depo.refresh_from_db(); self.assertFalse(self.depo.mutfak_var)
        # yeni şube mutfaklı açılır; depo + mutfak birlikte gelirse mutfak yok sayılır
        self.client.post(reverse('sube_yeni'), {'ad': 'Yeni', 'mutfak_var': 'on'})
        self.assertTrue(Sube.objects.get(ad='Yeni').mutfak_var)
        self.client.post(reverse('sube_yeni'), {'ad': 'Yeni Depo', 'depo_mu': 'on', 'mutfak_var': 'on'})
        self.assertFalse(Sube.objects.get(ad='Yeni Depo').mutfak_var)

    def test_yetkisiz_degistiremez(self):
        for rol, kw in ((Rol.MUDUR, {}), (Rol.SEF, dict(sube=self.c)), (Rol.YATIRIMCI, {}), (Rol.MUTFAK_SORUMLUSU, {})):
            self.gir(self.kisi(rol, **kw))
            self.client.post(reverse('sube_yeni'), {'islem': 'mutfak_ayar', 'sube_id': self.c.id, 'mutfak_var': 'on'})
        self.c.refresh_from_db(); self.assertFalse(self.c.mutfak_var)

    def test_kapatinca_eski_kayit_silinmez(self):
        self.kaydet(self.a, 'Domates', 1, 'kg')
        self.a.mutfak_var = False; self.a.save()
        self.assertEqual(MutfakZayiKalemi.objects.count(), 1)
        self.assertEqual(zayi_subeleri(self.kisi(Rol.GENEL_MUDUR)), [self.b])


@override_settings(**_AYAR)
class GocTestleri(ZayiBase):
    def test_mevcut_mutfakli_subeler_isaretlenir(self):
        import importlib
        goc = importlib.import_module('panel.migrations.0059_mutfak_zayi_yazili')
        Sube.objects.update(mutfak_var=False)
        mp = self.kisi(Rol.MUTFAK_PERSONEL, sube=self.a)
        Vardiya.objects.create(personel=mp, tarih=self.bugun, vardiya_tipi=VardiyaTipi.MUTFAK_GOREVI,
                               atanan_sube=self.b)
        MutfakZayi.objects.create(sube=self.depo, foto='x.jpg')                  # depo asla işaretlenmez
        goc.mevcut_mutfakli_subeleri_isaretle(global_apps, None)
        self.assertEqual(set(Sube.objects.filter(mutfak_var=True).values_list('ad', flat=True)), {'A Şube', 'B Şube'})


@override_settings(**_AYAR)
class ModelTestleri(ZayiBase):
    def test_kaydetme_alanlari(self):
        p = self.kisi(Rol.MUTFAK_PERSONEL, 'Ayşe', sube=self.a)
        k = MutfakZayiKalemi(sube=self.a, personel=p, tarih=self.bugun, urun='  PESTO   sos ', miktar=Decimal('1.5'), birim='lt')
        k.save()
        k.refresh_from_db()
        self.assertEqual((k.urun, k.urun_anahtar, k.olcu, k.taban_miktar), ('PESTO sos', 'pesto sos', 'hacim', Decimal('1500.00')))
        self.assertEqual((k.sube_ad_arsiv, k.personel_ad_arsiv), ('A Şube', 'Ayşe'))
        # şube / personel silinse de kayıt ve adlar kalır
        p.delete(); self.a.delete()
        k.refresh_from_db()
        self.assertIsNone(k.sube_id); self.assertEqual((k.sube_ad_arsiv, k.personel_ad_arsiv), ('A Şube', 'Ayşe'))
        self.assertEqual(zayi_rapor.sube_adi(k), 'A Şube')
