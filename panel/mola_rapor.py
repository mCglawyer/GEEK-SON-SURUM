"""Mola raporu — günlük toplam + hangi saatlerde çıkıldığı (Excel / PDF)."""
import io

from .mola_kurallari import MOLA_GUNLUK_LIMIT_DK


def _saatler(g):
    return ', '.join('%s–%s%s' % (m['giris'].strftime('%H:%M'), m['cikis'].strftime('%H:%M'),
                                  ' (oto)' if m['oto'] else '') for m in g['molalar'])


def excel_bytes(gruplar, etiket):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    wb = openpyxl.Workbook()
    mavi = PatternFill('solid', fgColor='162AA3')
    zebra = PatternFill('solid', fgColor='F4F6FB')
    kirmizi = PatternFill('solid', fgColor='FBE7E8')
    ince = Side(style='thin', color='D7DCEB')
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)

    def sayfa(ws, baslik, basliklar, satirlar, olculer, asim_kolonu=None):
        excel_logo(ws)
        ws['C2'] = baslik
        ws['C2'].font = Font(bold=True, size=14, color='162AA3')
        ws['C3'] = etiket
        ws['C3'].font = Font(color='6B7280')
        ust = 5
        for i, b in enumerate(basliklar, 1):
            c = ws.cell(row=ust, column=i, value=b)
            c.fill, c.font, c.border = mavi, Font(bold=True, color='FFFFFF'), kenar
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for r, (satir, asti) in enumerate(satirlar, ust + 1):
            for i, v in enumerate(satir, 1):
                c = ws.cell(row=r, column=i, value=v)
                c.border = kenar
                c.alignment = Alignment(vertical='top', wrap_text=True)
                if asti:
                    c.fill = kirmizi
                elif (r - ust) % 2 == 0:
                    c.fill = zebra
        for i, w in enumerate(olculer, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=ust + 1, column=1)

    ws = wb.active
    ws.title = 'Günlük Özet'
    sayfa(ws, 'Mola Raporu — Günlük Özet',
          ['Tarih', 'Personel', 'Şube', 'Mola sayısı', 'Toplam (dk)', 'Günlük hak (dk)', 'Durum', 'Çıkış–dönüş saatleri'],
          [([g['tarih'].strftime('%d.%m.%Y'), g['ad'], g['sube'], g['adet'], g['toplam_dk'], MOLA_GUNLUK_LIMIT_DK,
             'Aşım' if g['asti'] else 'Uygun', _saatler(g)], g['asti']) for g in gruplar],
          [12, 26, 22, 12, 12, 14, 10, 52])

    ws = wb.create_sheet('Tüm Molalar')
    satirlar = []
    for g in gruplar:
        for m in g['molalar']:
            satirlar.append(([g['tarih'].strftime('%d.%m.%Y'), g['ad'], g['sube'], m['giris'].strftime('%H:%M'),
                              m['cikis'].strftime('%H:%M'), m['dk'],
                              'Otomatik kapatıldı' if m['oto'] else ('Manuel' if m['manuel'] else '')], False))
    sayfa(ws, 'Mola Raporu — Tüm Molalar',
          ['Tarih', 'Personel', 'Şube', 'Çıkış', 'Dönüş', 'Süre (dk)', 'Not'], satirlar,
          [12, 26, 22, 10, 10, 12, 24])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_bytes(gruplar, etiket):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf
    from .zayi_rapor import _pdf_font

    FONT, FONTB = _pdf_font()
    BRAND = colors.HexColor('#162AA3')
    ZEBRA = colors.HexColor('#F4F6FB')
    LINE = colors.HexColor('#E3E6EF')
    MUTED = colors.HexColor('#6B7280')
    ASIM = colors.HexColor('#FBE7E8')
    st_baslik = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_alt = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=8, leading=10.2)
    st_cw = ParagraphStyle('cw', fontName=FONTB, fontSize=8, leading=10.2, textColor=colors.white)
    st_not = ParagraphStyle('n', fontName=FONT, fontSize=7.6, textColor=MUTED, leading=10)
    LEFT = RIGHT = 15 * mm
    W = A4[0] - LEFT - RIGHT

    def P(t, st=st_c):
        return Paragraph(escape(str(t)), st)

    el = [Paragraph('Mola Raporu', st_baslik),
          Paragraph(escape(etiket), st_alt),
          Paragraph('Günlük hak: %d dk · Kayıt: %d gün/kişi' % (MOLA_GUNLUK_LIMIT_DK, len(gruplar)), st_alt),
          Spacer(1, 8)]
    rows = [[P(h, st_cw) for h in ('Tarih', 'Personel', 'Şube', 'Mola', 'Toplam', 'Çıkış–dönüş saatleri')]]
    asimlar = []
    for i, g in enumerate(gruplar, 1):
        rows.append([P(g['tarih'].strftime('%d.%m.%Y')), P(g['ad']), P(g['sube']), P(g['adet']),
                     P('%d dk%s' % (g['toplam_dk'], ' ⚠' if g['asti'] else '')), P(_saatler(g))])
        if g['asti']:
            asimlar.append(i)
    if len(rows) == 1:
        el.append(Paragraph('Bu aralıkta mola kaydı yok.', st_alt))
    else:
        t = Table(rows, colWidths=[W * .11, W * .21, W * .16, W * .07, W * .10, W * .35], repeatRows=1)
        stil = [('BACKGROUND', (0, 0), (-1, 0), BRAND), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, ZEBRA]),
                ('GRID', (0, 0), (-1, -1), 0.3, LINE), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4)]
        stil += [('BACKGROUND', (0, i), (-1, i), ASIM) for i in asimlar]
        t.setStyle(TableStyle(stil))
        el.append(t)
    el.append(Spacer(1, 8))
    el.append(Paragraph('Toplam: o gün yapılan tüm molaların toplamı. (oto): dönüş QR’ı okutulmadığı için otomatik kapatılan mola.', st_not))
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=A4, left_margin=LEFT, right_margin=RIGHT, font=FONT, title='Mola Raporu')
    return buf.getvalue()
