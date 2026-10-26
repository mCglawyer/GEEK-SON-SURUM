"""Satış raporu (Menulux) Excel / PDF çıktıları."""
import io
from xml.sax.saxutils import escape

from . import menulux as mx


def tl(v):
    s = '{:,.2f}'.format(float(v or 0)).replace(',', 'X').replace('.', ',').replace('X', '.')
    return s


def adet(v):
    v = float(v or 0)
    if v == int(v):
        return '{:,}'.format(int(v)).replace(',', '.')
    return '{:,.1f}'.format(v).replace(',', 'X').replace('.', ',').replace('X', '.')


def _ozet_satirlari(o):
    return [('Ciro (Menulux toplam)', tl(o['ciro'])), ('Brüt satış', tl(o['brut'])), ('İndirim', tl(o['indirim'])),
            ('İkram', tl(o['ikram'])), ('İptal', tl(o['iptal'])), ('İade', tl(o['iade'])), ('KDV', tl(o['kdv'])),
            ('Adisyon sayısı', adet(o['adisyon'])), ('Kişi sayısı', adet(o['kisi'])), ('Ortalama adisyon', tl(mx.ortalama(o)))]


def excel_bytes(veri, etiket, sube_etiket):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    mavi = PatternFill('solid', fgColor='162AA3'); kenar = Border(*[Side(style='thin', color='D7DCEB')] * 4)
    bas_font = Font(bold=True, color='FFFFFF'); sayi_bicim = '#,##0.00'; adet_bicim = '#,##0.##'
    wb = Workbook()

    def sayfa(ws, baslik, kolonlar, satirlar, bicimler=None, genislik=None):
        ws['A1'] = baslik; ws['A1'].font = Font(bold=True, size=14, color='162AA3')
        ws['A2'] = '%s · %s' % (sube_etiket, etiket); ws['A2'].font = Font(italic=True, color='5B6472')
        for j, k in enumerate(kolonlar, 1):
            c = ws.cell(row=4, column=j, value=k); c.fill, c.font, c.border = mavi, bas_font, kenar
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i, s in enumerate(satirlar, 5):
            for j, v in enumerate(s, 1):
                c = ws.cell(row=i, column=j, value=v); c.border = kenar
                if bicimler and bicimler[j - 1] and isinstance(v, (int, float)):
                    c.number_format = bicimler[j - 1]
        for j in range(1, len(kolonlar) + 1):
            ws.column_dimensions[get_column_letter(j)].width = (genislik or [30] + [16] * 12)[j - 1] if j <= len(genislik or [30] + [16] * 12) else 16
        ws.freeze_panes = 'A5'
        excel_logo(ws)

    o = veri['ozet']
    ws = wb.active; ws.title = 'Özet'
    satirlar = [(a, float(b.replace('.', '').replace(',', '.'))) for a, b in _ozet_satirlari(o)]
    satirlar += [('Ödeme: ' + ad, v) for ad, v in o['odeme']]
    sayfa(ws, 'Satış özeti', ['Kalem', 'Değer'], satirlar, [None, sayi_bicim], [34, 18])

    if veri['subeler']:
        sayfa(wb.create_sheet('Şubeler'), 'Şube bazlı satış',
              ['Şube', 'Ciro', 'Adisyon', 'Kişi', 'Ort. adisyon', 'İndirim', 'İkram', 'İptal', 'İade', 'KDV'],
              [(s['ad'], s['ciro'], s['adisyon'], s['kisi'], s['ort'], s['indirim'], s['ikram'], s['iptal'], s['iade'], s['kdv'])
               for s in veri['subeler']],
              [None, sayi_bicim, '#,##0', '#,##0', sayi_bicim, sayi_bicim, sayi_bicim, sayi_bicim, sayi_bicim, sayi_bicim], [34] + [15] * 9)
    if veri['gunluk']:
        sayfa(wb.create_sheet('Günlük'), 'Günlük satış', ['Tarih', 'Ciro', 'Adisyon', 'Kişi', 'Ort. adisyon'],
              [(g.strftime('%d.%m.%Y'), x['ciro'], x['adisyon'], x['kisi'], mx.ortalama(x)) for g, x in veri['gunluk']],
              [None, sayi_bicim, '#,##0', '#,##0', sayi_bicim], [16, 16, 12, 12, 16])
    if veri['kategoriler']:
        sayfa(wb.create_sheet('Kategoriler'), 'Kategori satışları', ['Kategori', 'Adet', 'Net', 'KDV', 'Toplam'],
              [(k['ad'], k['adet'], k['net'], k['kdv'], k['tutar']) for k in veri['kategoriler']],
              [None, adet_bicim, sayi_bicim, sayi_bicim, sayi_bicim], [34, 12, 16, 16, 16])
    if 'urunler' in veri:
        sayfa(wb.create_sheet('Satan ürünler'), 'Satan ürünler', ['Ürün', 'Kod', 'Grup', 'Adet', 'Tutar', 'KDV'],
              [(u['ad'], u['kod'], u['grup'], u['adet'], u['tutar'], u['kdv']) for u in veri['urunler']],
              [None, None, None, adet_bicim, sayi_bicim, sayi_bicim], [38, 12, 26, 12, 16, 14])
        if veri.get('satilmayan_hata'):
            sayfa(wb.create_sheet('Satılmayan'), 'Satılmayan ürünler', ['Not'], [(veri['satilmayan_hata'],)], None, [80])
        else:
            sayfa(wb.create_sheet('Satılmayan'), 'Satılmayan ürünler', ['Ürün', 'Kod', 'Grup', 'Liste fiyatı'],
                  [(u['ad'], u['kod'], u['grup'], u['fiyat']) for u in veri['satilmayan']], [None, None, None, sayi_bicim], [38, 12, 26, 14])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def pdf_bytes(veri, etiket, sube_etiket):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf
    from .zayi_rapor import _pdf_font

    FONT, FONTB = _pdf_font()
    BRAND = colors.HexColor('#162AA3'); MUTED = colors.HexColor('#5b6472'); LINE = colors.HexColor('#d7dceb'); SOFT = colors.HexColor('#EEF1FB')
    st_b = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_a = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_s = ParagraphStyle('s', fontName=FONTB, fontSize=11, textColor=BRAND, leading=15, spaceBefore=8)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=8, leading=10)
    st_n = ParagraphStyle('n', fontName=FONT, fontSize=8, leading=10, textColor=MUTED)
    P = lambda t, s=st_c: Paragraph(escape(str(t)), s)

    def tablo(kolonlar, satirlar, genislik, sag=()):
        data = [[P(k, ParagraphStyle('h', fontName=FONTB, fontSize=8, leading=10, textColor=colors.white)) for k in kolonlar]]
        for s in satirlar:
            data.append([P(v) for v in s])
        t = Table(data, colWidths=[g * mm for g in genislik], repeatRows=1)
        stil = [('BACKGROUND', (0, 0), (-1, 0), BRAND), ('GRID', (0, 0), (-1, -1), 0.4, LINE), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, SOFT]), ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
        for c in sag:
            stil.append(('ALIGN', (c, 1), (c, -1), 'RIGHT'))
        t.setStyle(TableStyle(stil)); return t

    o = veri['ozet']
    el = [P('Satış raporu', st_b), P('%s · %s' % (sube_etiket, etiket), st_a), Spacer(1, 4 * mm)]
    el.append(tablo(['Kalem', 'Değer'], _ozet_satirlari(o) + [('Ödeme: ' + a, tl(v)) for a, v in o['odeme']], [90, 40], sag=(1,)))
    if veri['subeler']:
        el += [P('Şube bazlı satış', st_s), tablo(['Şube', 'Ciro', 'Adisyon', 'Ort. adisyon'],
              [(s['ad'], tl(s['ciro']), adet(s['adisyon']), tl(s['ort'])) for s in veri['subeler']], [80, 32, 24, 32], sag=(1, 2, 3))]
    if veri['gunluk']:
        el += [P('Günlük satış', st_s), tablo(['Tarih', 'Ciro', 'Adisyon', 'Ort. adisyon'],
              [(g.strftime('%d.%m.%Y'), tl(x['ciro']), adet(x['adisyon']), tl(mx.ortalama(x))) for g, x in veri['gunluk']], [40, 40, 30, 40], sag=(1, 2, 3))]
    if veri['kategoriler']:
        el += [P('Kategori satışları', st_s), tablo(['Kategori', 'Adet', 'Toplam'],
              [(k['ad'], adet(k['adet']), tl(k['tutar'])) for k in veri['kategoriler']], [90, 30, 40], sag=(1, 2))]
    if 'urunler' in veri:
        el += [P('Satan ürünler', st_s), tablo(['Ürün', 'Grup', 'Adet', 'Tutar'],
              [(u['ad'], u['grup'], adet(u['adet']), tl(u['tutar'])) for u in veri['urunler']], [66, 50, 22, 34], sag=(2, 3))]
        el.append(P('Satılmayan ürünler', st_s))
        if veri.get('satilmayan_hata'):
            el.append(P(veri['satilmayan_hata'], st_n))
        elif veri['satilmayan']:
            el.append(tablo(['Ürün', 'Grup', 'Liste fiyatı'], [(u['ad'], u['grup'], tl(u['fiyat'])) for u in veri['satilmayan']], [80, 60, 32], sag=(2,)))
        else:
            el.append(P('Bu aralıkta satılmayan ürün yok.', st_n))
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=A4, left_margin=14 * mm, right_margin=14 * mm, font=FONT, title='Satış raporu')
    return buf.getvalue()


# ------------------------------------------------------------------ Rapor Özeti çıktıları
def ozet_excel_bytes(veri, etiket, sube_etiket):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    mavi = PatternFill('solid', fgColor='162AA3'); kenar = Border(*[Side(style='thin', color='D7DCEB')] * 4)
    bas_font = Font(bold=True, color='FFFFFF'); sayi_bicim = '#,##0.00'; adet_bicim = '#,##0.##'
    wb = Workbook()

    def sayfa(ws, baslik, kolonlar, satirlar, bicimler=None, genislik=None):
        ws['A1'] = baslik; ws['A1'].font = Font(bold=True, size=14, color='162AA3')
        ws['A2'] = '%s · %s' % (sube_etiket, etiket); ws['A2'].font = Font(italic=True, color='5B6472')
        for j, k in enumerate(kolonlar, 1):
            c = ws.cell(row=4, column=j, value=k); c.fill, c.font, c.border = mavi, bas_font, kenar
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i, s in enumerate(satirlar, 5):
            for j, v in enumerate(s, 1):
                c = ws.cell(row=i, column=j, value=v); c.border = kenar
                if bicimler and bicimler[j - 1] and isinstance(v, (int, float)):
                    c.number_format = bicimler[j - 1]
        for j in range(1, len(kolonlar) + 1):
            ws.column_dimensions[get_column_letter(j)].width = (genislik or [30] + [16] * 12)[j - 1] if j <= len(genislik or [30] + [16] * 12) else 16
        ws.freeze_panes = 'A5'
        excel_logo(ws)

    o = veri['ozet']
    ws = wb.active; ws.title = 'Özet'
    satirlar = [(a, float(b.replace('.', '').replace(',', '.'))) for a, b in _ozet_satirlari(o)]
    satirlar += [('Ödeme: ' + ad, v) for ad, v in o.get('odeme') or []]
    sayfa(ws, 'Rapor özeti', ['Kalem', 'Değer'], satirlar, [None, sayi_bicim], [34, 18])

    if veri.get('grafik'):
        sayfa(wb.create_sheet('Günlük ciro'), 'Günlük ciro', ['Tarih', 'Ciro'], list(veri['grafik']), [None, sayi_bicim], [18, 18])
    if veri.get('subeler'):
        sayfa(wb.create_sheet('Şubeler'), 'Şube bazlı satış',
              ['Şube', 'Ciro', 'Adisyon', 'Kişi', 'Ort. adisyon', 'İndirim', 'İkram', 'İptal', 'İade', 'KDV'],
              [(s['ad'], s['ciro'], s['adisyon'], s['kisi'], s['ort'], s['indirim'], s['ikram'], s['iptal'], s['iade'], s['kdv'])
               for s in veri['subeler']],
              [None, sayi_bicim, '#,##0', '#,##0', sayi_bicim, sayi_bicim, sayi_bicim, sayi_bicim, sayi_bicim, sayi_bicim], [34] + [15] * 9)
    if veri.get('kategoriler'):
        sayfa(wb.create_sheet('Kategoriler'), 'Kategori satışları', ['Kategori', 'Adet', 'Net', 'KDV', 'Toplam'],
              [(k['ad'], k['adet'], k['net'], k['kdv'], k['tutar']) for k in veri['kategoriler']],
              [None, adet_bicim, sayi_bicim, sayi_bicim, sayi_bicim], [34, 12, 16, 16, 16])
    if veri.get('saatlik'):
        sayfa(wb.create_sheet('Saat bazlı kanal'), 'Saat bazlı kasa/paket', ['Saat', 'Kasa', 'Paket'],
              [('%02d:00' % s['saat'], s['kasa'], s['paket']) for s in veri['saatlik']], [None, sayi_bicim, sayi_bicim], [12, 16, 16])
    if veri.get('noktalar'):
        sayfa(wb.create_sheet('Satış noktaları'), 'Satış noktaları', ['Nokta', 'Tutar', 'Pay (%)'],
              [(n['ad'], n['tutar'], n['pay']) for n in veri['noktalar']], [None, sayi_bicim, '#,##0.0'], [34, 18, 14])
    if veri.get('urunler'):
        sayfa(wb.create_sheet('Satan ürünler'), 'En çok satan ürünler', ['Ürün', 'Kod', 'Grup', 'Adet', 'Tutar', 'KDV'],
              [(u['ad'], u['kod'], u['grup'], u['adet'], u['tutar'], u['kdv']) for u in veri['urunler'][:50]],
              [None, None, None, adet_bicim, sayi_bicim, sayi_bicim], [38, 12, 26, 12, 16, 14])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def ozet_pdf_bytes(veri, etiket, sube_etiket):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf
    from .zayi_rapor import _pdf_font

    FONT, FONTB = _pdf_font()
    BRAND = colors.HexColor('#162AA3'); MUTED = colors.HexColor('#5b6472'); LINE = colors.HexColor('#d7dceb'); SOFT = colors.HexColor('#EEF1FB')
    st_b = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_a = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_s = ParagraphStyle('s', fontName=FONTB, fontSize=11, textColor=BRAND, leading=15, spaceBefore=8)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=8, leading=10)
    P = lambda t, s=st_c: Paragraph(escape(str(t)), s)

    def tablo(kolonlar, satirlar, genislik, sag=()):
        data = [[P(k, ParagraphStyle('h', fontName=FONTB, fontSize=8, leading=10, textColor=colors.white)) for k in kolonlar]]
        for s in satirlar:
            data.append([P(v) for v in s])
        t = Table(data, colWidths=[g * mm for g in genislik], repeatRows=1)
        stil = [('BACKGROUND', (0, 0), (-1, 0), BRAND), ('GRID', (0, 0), (-1, -1), 0.4, LINE), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, SOFT]), ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
        for c in sag:
            stil.append(('ALIGN', (c, 1), (c, -1), 'RIGHT'))
        t.setStyle(TableStyle(stil)); return t

    o = veri['ozet']
    el = [P('Rapor özeti', st_b), P('%s · %s' % (sube_etiket, etiket), st_a), Spacer(1, 4 * mm)]
    el.append(tablo(['Kalem', 'Değer'], _ozet_satirlari(o) + [('Ödeme: ' + a, tl(v)) for a, v in o.get('odeme') or []], [90, 40], sag=(1,)))
    if veri.get('grafik'):
        el += [P('Günlük ciro', st_s), tablo(['Tarih', 'Ciro'], [(g, tl(v)) for g, v in veri['grafik']], [70, 40], sag=(1,))]
    if veri.get('subeler'):
        el += [P('Şube bazlı satış', st_s), tablo(['Şube', 'Ciro', 'Adisyon', 'Ort. adisyon'],
              [(s['ad'], tl(s['ciro']), adet(s['adisyon']), tl(s['ort'])) for s in veri['subeler']], [80, 32, 24, 32], sag=(1, 2, 3))]
    if veri.get('kategoriler'):
        el += [P('Kategori satışları', st_s), tablo(['Kategori', 'Adet', 'Toplam'],
              [(k['ad'], adet(k['adet']), tl(k['tutar'])) for k in veri['kategoriler']], [90, 30, 40], sag=(1, 2))]
    if veri.get('saatlik'):
        el += [P('Saat bazlı kanal (kasa / paket)', st_s), tablo(['Saat', 'Kasa', 'Paket'],
              [('%02d:00' % s['saat'], tl(s['kasa']), tl(s['paket'])) for s in veri['saatlik']], [30, 45, 45], sag=(1, 2))]
    if veri.get('noktalar'):
        el += [P('Satış noktaları', st_s), tablo(['Nokta', 'Tutar', 'Pay'],
              [(n['ad'], tl(n['tutar']), '%' + ('%.1f' % n['pay'])) for n in veri['noktalar']], [80, 35, 25], sag=(1, 2))]
    if veri.get('urunler'):
        el += [P('En çok satan ürünler', st_s), tablo(['Ürün', 'Grup', 'Adet', 'Tutar'],
              [(u['ad'], u['grup'], adet(u['adet']), tl(u['tutar'])) for u in veri['urunler'][:25]], [66, 50, 22, 34], sag=(2, 3))]
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=A4, left_margin=14 * mm, right_margin=14 * mm, font=FONT, title='Rapor özeti')
    return buf.getvalue()


# ------------------------------------------------------------------ İkram / iptal / iade detayı çıktıları
def _ii_kolonlar(veri, basliklar):
    kolonlar = ['Zaman']
    cokSube = any(s.get('sube') for liste in veri.values() for s in liste)
    if cokSube:
        kolonlar.append('Şube')
    kolonlar += ['Ürün', 'Adet', 'Liste tutarı', 'Masa/Kasa', 'Personel', 'Açıklama']
    return kolonlar, cokSube


def ikram_iptal_excel_bytes(veri, etiket, sube_etiket):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    mavi = PatternFill('solid', fgColor='162AA3'); kenar = Border(*[Side(style='thin', color='D7DCEB')] * 4)
    bas_font = Font(bold=True, color='FFFFFF')
    wb = Workbook()
    kolonlar, cokSube = _ii_kolonlar(veri, None)

    def sayfa(ws, baslik, satirlar):
        ws['A1'] = baslik; ws['A1'].font = Font(bold=True, size=14, color='162AA3')
        ws['A2'] = '%s · %s' % (sube_etiket, etiket); ws['A2'].font = Font(italic=True, color='5B6472')
        for j, k in enumerate(kolonlar, 1):
            c = ws.cell(row=4, column=j, value=k); c.fill, c.font, c.border = mavi, bas_font, kenar
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i, s in enumerate(satirlar, 5):
            satir = [s['zaman'][:16].replace('T', ' ')]
            if cokSube:
                satir.append(s['sube'])
            satir += [s['urun'], s['adet'], s['tutar'], s['masa'], s['personel'], s['aciklama']]
            for j, v in enumerate(satir, 1):
                c = ws.cell(row=i, column=j, value=v); c.border = kenar
        genislik = [16] + ([22] if cokSube else []) + [26, 10, 14, 14, 20, 40]
        for j in range(1, len(kolonlar) + 1):
            ws.column_dimensions[get_column_letter(j)].width = genislik[j - 1] if j <= len(genislik) else 16
        ws.freeze_panes = 'A5'
        excel_logo(ws)

    ws = wb.active; ws.title = 'İkram'; sayfa(ws, 'İkram edilen ürünler', veri.get('ikram') or [])
    sayfa(wb.create_sheet('İptal'), 'İptal edilen ürünler', veri.get('iptal') or [])
    sayfa(wb.create_sheet('İade'), 'İade edilen ürünler', veri.get('iade') or [])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def ikram_iptal_pdf_bytes(veri, etiket, sube_etiket):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf
    from .zayi_rapor import _pdf_font

    FONT, FONTB = _pdf_font()
    BRAND = colors.HexColor('#162AA3'); MUTED = colors.HexColor('#5b6472'); LINE = colors.HexColor('#d7dceb'); SOFT = colors.HexColor('#EEF1FB')
    st_b = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_a = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_s = ParagraphStyle('s', fontName=FONTB, fontSize=11, textColor=BRAND, leading=15, spaceBefore=8)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=7.5, leading=9.5)
    st_n = ParagraphStyle('n', fontName=FONT, fontSize=8, leading=10, textColor=MUTED)
    P = lambda t, s=st_c: Paragraph(escape(str(t)), s)

    kolonlar, cokSube = _ii_kolonlar(veri, None)

    def tablo(satirlar):
        data = [[P(k, ParagraphStyle('h', fontName=FONTB, fontSize=7.5, leading=9.5, textColor=colors.white)) for k in kolonlar]]
        for s in satirlar:
            satir = [s['zaman'][:16].replace('T', ' ')]
            if cokSube:
                satir.append(s['sube'])
            satir += [s['urun'], adet(s['adet']), tl(s['tutar']), s['masa'] or '—', s['personel'] or '—', s['aciklama'] or '—']
            data.append([P(v) for v in satir])
        genislik = [22] + ([28] if cokSube else []) + [34, 14, 18, 18, 24] + [38 if cokSube else 46]
        t = Table(data, colWidths=[g * mm for g in genislik], repeatRows=1)
        stil = [('BACKGROUND', (0, 0), (-1, 0), BRAND), ('GRID', (0, 0), (-1, -1), 0.4, LINE), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, SOFT]), ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
        t.setStyle(TableStyle(stil)); return t

    el = [P('İkram & iptal detayı', st_b), P('%s · %s' % (sube_etiket, etiket), st_a), Spacer(1, 4 * mm)]
    for baslik, anahtar in (('İkram edilen ürünler', 'ikram'), ('İptal edilen ürünler', 'iptal'), ('İade edilen ürünler', 'iade')):
        el.append(P(baslik, st_s))
        liste = veri.get(anahtar) or []
        el.append(tablo(liste) if liste else P('Bu aralıkta kayıt yok.', st_n))
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=A4, left_margin=12 * mm, right_margin=12 * mm, font=FONT, title='İkram & iptal detayı')
    return buf.getvalue()


# ------------------------------------------------------------------ Tablo görünümü çıktıları
def tablo_excel_bytes(ctx, etiket, sube_etiket):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    mavi = PatternFill('solid', fgColor='162AA3'); kenar = Border(*[Side(style='thin', color='D7DCEB')] * 4)
    bas_font = Font(bold=True, color='FFFFFF'); sayi_bicim = '#,##0.00'; adet_bicim = '#,##0.##'
    wb = Workbook()

    def sayfa(ws, baslik, kolonlar, satirlar, bicimler=None, genislik=None):
        ws['A1'] = baslik; ws['A1'].font = Font(bold=True, size=14, color='162AA3')
        ws['A2'] = '%s · %s' % (sube_etiket, etiket); ws['A2'].font = Font(italic=True, color='5B6472')
        for j, k in enumerate(kolonlar, 1):
            c = ws.cell(row=4, column=j, value=k); c.fill, c.font, c.border = mavi, bas_font, kenar
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i, s in enumerate(satirlar, 5):
            for j, v in enumerate(s, 1):
                c = ws.cell(row=i, column=j, value=v); c.border = kenar
                if bicimler and bicimler[j - 1] and isinstance(v, (int, float)):
                    c.number_format = bicimler[j - 1]
        for j in range(1, len(kolonlar) + 1):
            ws.column_dimensions[get_column_letter(j)].width = (genislik or [30] + [16] * 12)[j - 1] if j <= len(genislik or [30] + [16] * 12) else 16
        ws.freeze_panes = 'A5'
        excel_logo(ws)

    ws = wb.active; ws.title = 'Dönemler'
    sayfa(ws, 'Dönem karşılaştırması', ['Dönem', 'Ciro', 'Adisyon', 'Kişi', 'Ort. adisyon', 'İndirim', 'İptal'],
          [(d['etiket'], d['ozet']['ciro'], d['ozet']['adisyon'], d['ozet']['kisi'], d['ozet'].get('ort', 0), d['ozet']['indirim'], d['ozet']['iptal'])
           for d in ctx['donem_satirlari']],
          [None, sayi_bicim, '#,##0', '#,##0', sayi_bicim, sayi_bicim, sayi_bicim], [18, 16, 12, 12, 16, 14, 14])
    if ctx.get('kategori'):
        sayfa(wb.create_sheet('Kategoriler'), 'Kategori satışları (%s)' % ctx.get('donem_etiket', ''), ['Kategori', 'Adet', 'Tutar', 'Pay (%)'],
              [(k['ad'], k['adet'], k['tutar'], k.get('pay', 0)) for k in ctx['kategori']],
              [None, adet_bicim, sayi_bicim, '#,##0.0'], [34, 12, 16, 14])
    if ctx.get('subeler'):
        sayfa(wb.create_sheet('Şubeler'), 'Şube karşılaştırması (%s)' % ctx.get('donem_etiket', ''),
              ['Şube', 'Ciro', 'Adisyon', 'Ort. adisyon', 'İndirim', 'İptal'],
              [(s['ad'], s['ciro'], s['adisyon'], s['ort'], s['indirim'], s['iptal']) for s in ctx['subeler']],
              [None, sayi_bicim, '#,##0', sayi_bicim, sayi_bicim, sayi_bicim], [34, 16, 12, 16, 14, 14])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def tablo_pdf_bytes(ctx, etiket, sube_etiket):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf
    from .zayi_rapor import _pdf_font

    FONT, FONTB = _pdf_font()
    BRAND = colors.HexColor('#162AA3'); MUTED = colors.HexColor('#5b6472'); LINE = colors.HexColor('#d7dceb'); SOFT = colors.HexColor('#EEF1FB')
    st_b = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_a = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_s = ParagraphStyle('s', fontName=FONTB, fontSize=11, textColor=BRAND, leading=15, spaceBefore=8)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=8, leading=10)
    P = lambda t, s=st_c: Paragraph(escape(str(t)), s)

    def tablo(kolonlar, satirlar, genislik, sag=()):
        data = [[P(k, ParagraphStyle('h', fontName=FONTB, fontSize=8, leading=10, textColor=colors.white)) for k in kolonlar]]
        for s in satirlar:
            data.append([P(v) for v in s])
        t = Table(data, colWidths=[g * mm for g in genislik], repeatRows=1)
        stil = [('BACKGROUND', (0, 0), (-1, 0), BRAND), ('GRID', (0, 0), (-1, -1), 0.4, LINE), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, SOFT]), ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
        for c in sag:
            stil.append(('ALIGN', (c, 1), (c, -1), 'RIGHT'))
        t.setStyle(TableStyle(stil)); return t

    el = [P('Tablo görünümü', st_b), P('%s · %s' % (sube_etiket, etiket), st_a), Spacer(1, 4 * mm)]
    el.append(tablo(['Dönem', 'Ciro', 'Adisyon', 'Ort. adisyon', 'İptal'],
              [(d['etiket'], tl(d['ozet']['ciro']), adet(d['ozet']['adisyon']), tl(d['ozet'].get('ort', 0)), tl(d['ozet']['iptal']))
               for d in ctx['donem_satirlari']], [34, 32, 26, 32, 28], sag=(1, 2, 3, 4)))
    if ctx.get('kategori'):
        el += [P('Kategori satışları (%s)' % ctx.get('donem_etiket', ''), st_s), tablo(['Kategori', 'Adet', 'Tutar'],
              [(k['ad'], adet(k['adet']), tl(k['tutar'])) for k in ctx['kategori']], [90, 30, 40], sag=(1, 2))]
    if ctx.get('subeler'):
        el += [P('Şube karşılaştırması (%s)' % ctx.get('donem_etiket', ''), st_s), tablo(['Şube', 'Ciro', 'Adisyon', 'Ort. adisyon'],
              [(s['ad'], tl(s['ciro']), adet(s['adisyon']), tl(s['ort'])) for s in ctx['subeler']], [80, 32, 24, 32], sag=(1, 2, 3))]
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=A4, left_margin=14 * mm, right_margin=14 * mm, font=FONT, title='Tablo görünümü')
    return buf.getvalue()
