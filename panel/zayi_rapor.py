"""Mutfak zayi raporu — dönem hesabı, toplamlar, Excel ve PDF çıktısı.

Toplama kuralı: ağırlık (gr+kg), hacim (ml+lt) ve adet AYRI toplanır; gram ile
mililitre ya da adet asla birbirine eklenmez. Ürün adları büyük/küçük harf ve
fazladan boşluk farkına bakılmadan aynı ürün sayılır (Domates = domates).
"""
import datetime
import io
import os
from decimal import Decimal

from django.db.models import Count, Sum

from .models import MutfakZayiKalemi, zayi_urun_baslik

DONEMLER = (('gun', 'Gün'), ('hafta', 'Hafta'), ('ay', 'Ay'), ('ozel', 'Tarih aralığı'))
OZEL_MAKS_GUN = 366

OLCULER = ('agirlik', 'hacim', 'adet')
OLCU_ETIKET = {'agirlik': 'Ağırlık', 'hacim': 'Hacim', 'adet': 'Adet'}
AYLAR = ['Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos',
         'Eylül', 'Ekim', 'Kasım', 'Aralık']


# ---- Dönem ------------------------------------------------------------------------
def _ay_sonu(d):
    ilk = d.replace(day=1)
    sonraki = (ilk.replace(day=28) + datetime.timedelta(days=4)).replace(day=1)
    return sonraki - datetime.timedelta(days=1)


def tarih_oku(metin, varsayilan):
    try:
        return datetime.datetime.strptime((metin or '').strip(), '%Y-%m-%d').date()
    except ValueError:
        return varsayilan


def donem_hesapla(donem, tarih, bas=None, bit=None):
    """-> (donem, başlangıç, bitiş, etiket, önceki_tarih, sonraki_tarih, uyarı)"""
    if donem not in dict(DONEMLER):
        donem = 'gun'
    uyari = ''
    if donem == 'gun':
        b = e = tarih
        etiket = tarih.strftime('%d.%m.%Y')
        onceki, sonraki = tarih - datetime.timedelta(days=1), tarih + datetime.timedelta(days=1)
    elif donem == 'hafta':
        b = tarih - datetime.timedelta(days=tarih.weekday())
        e = b + datetime.timedelta(days=6)
        etiket = '%s – %s' % (b.strftime('%d.%m.%Y'), e.strftime('%d.%m.%Y'))
        onceki, sonraki = b - datetime.timedelta(days=7), b + datetime.timedelta(days=7)
    elif donem == 'ay':
        b = tarih.replace(day=1)
        e = _ay_sonu(b)
        etiket = '%s %d' % (AYLAR[b.month - 1], b.year)
        onceki, sonraki = (b - datetime.timedelta(days=1)).replace(day=1), e + datetime.timedelta(days=1)
    else:
        b = bas or tarih
        e = bit or tarih
        if e < b:
            b, e = e, b
        if (e - b).days + 1 > OZEL_MAKS_GUN:
            e = b + datetime.timedelta(days=OZEL_MAKS_GUN - 1)
            uyari = 'En fazla %d günlük aralık seçilebilir; bitiş tarihi kısaltıldı.' % OZEL_MAKS_GUN
        etiket = '%s – %s' % (b.strftime('%d.%m.%Y'), e.strftime('%d.%m.%Y'))
        onceki = sonraki = None
    return donem, b, e, etiket, onceki, sonraki, uyari


# ---- Biçim ------------------------------------------------------------------------
def _sayi(d, ondalik=2):
    d = Decimal(d or 0)
    metin = ('%.' + str(ondalik) + 'f') % d
    if '.' in metin:
        metin = metin.rstrip('0').rstrip('.')
    return metin.replace('.', ',')


def miktar_yaz(taban, olcu):
    """Taban birimden (gr / ml / adet) okunaklı metin: 1250 gr -> '1,25 kg'."""
    taban = Decimal(taban or 0)
    if olcu == 'agirlik':
        return ('%s kg' % _sayi(taban / 1000)) if taban >= 1000 else ('%s gr' % _sayi(taban))
    if olcu == 'hacim':
        return ('%s lt' % _sayi(taban / 1000)) if taban >= 1000 else ('%s ml' % _sayi(taban))
    return '%s adet' % _sayi(taban)


def kalem_miktar_yaz(k):
    return '%s %s' % (_sayi(k.miktar), k.birim)


# ---- Veri -------------------------------------------------------------------------
def kalemler(sube_idleri, baslangic, bitis):
    return (MutfakZayiKalemi.objects
            .filter(sube_id__in=list(sube_idleri), tarih__gte=baslangic, tarih__lte=bitis)
            .select_related('sube', 'personel'))


def _bos():
    return {o: Decimal(0) for o in OLCULER}


def _toplam_yaz(t):
    return {o: miktar_yaz(t[o], o) if t[o] else '' for o in OLCULER}


def rapor_verisi(subeler, baslangic, bitis):
    """subeler: Sube listesi (kapsamdaki). Döner: sözlük (şablon + Excel + PDF ortak)."""
    sube_map = {s.id: s for s in subeler}
    qs = kalemler(sube_map.keys(), baslangic, bitis)

    # Şube özeti (verisi olmayan şubeler de 0 ile listelenir)
    ozet = {s.id: {'sube': s, 'kayit': 0, 'top': _bos()} for s in subeler}
    for r in qs.values('sube_id', 'olcu').annotate(t=Sum('taban_miktar'), n=Count('id')):
        o = ozet[r['sube_id']]
        o['top'][r['olcu']] += r['t'] or 0
        o['kayit'] += r['n']
    genel = {'kayit': 0, 'top': _bos()}
    ozet_liste = []
    for s in subeler:
        o = ozet[s.id]
        o['yazi'] = _toplam_yaz(o['top'])
        genel['kayit'] += o['kayit']
        for k in OLCULER:
            genel['top'][k] += o['top'][k]
        ozet_liste.append(o)
    genel['yazi'] = _toplam_yaz(genel['top'])

    # Ürün x şube (aynı ürün farklı ölçüde girildiyse ayrı satır: "Süt" lt ve adet gibi)
    urun_map = {}
    for r in qs.values('urun_anahtar', 'olcu', 'sube_id').annotate(t=Sum('taban_miktar')):
        anahtar = (r['urun_anahtar'], r['olcu'])
        u = urun_map.setdefault(anahtar, {
            'ad': zayi_urun_baslik(r['urun_anahtar']), 'olcu': r['olcu'],
            'subeler': {}, 'toplam': Decimal(0)})
        u['subeler'][r['sube_id']] = r['t'] or Decimal(0)
        u['toplam'] += r['t'] or Decimal(0)
    urunler = sorted(urun_map.values(), key=lambda u: (OLCULER.index(u['olcu']), -u['toplam'], u['ad']))
    for u in urunler:
        u['toplam_yazi'] = miktar_yaz(u['toplam'], u['olcu'])
        u['hucreler'] = [
            {'sube': s, 'taban': u['subeler'].get(s.id), 'yazi': miktar_yaz(u['subeler'][s.id], u['olcu'])
             if u['subeler'].get(s.id) else ''} for s in subeler]

    # Gün gün özet
    gun_map = {}
    for r in qs.values('tarih', 'olcu').annotate(t=Sum('taban_miktar'), n=Count('id')):
        g = gun_map.setdefault(r['tarih'], {'tarih': r['tarih'], 'kayit': 0, 'top': _bos()})
        g['top'][r['olcu']] += r['t'] or 0
        g['kayit'] += r['n']
    gunler = [gun_map[t] for t in sorted(gun_map)]
    for g in gunler:
        g['yazi'] = _toplam_yaz(g['top'])

    detay = list(qs.order_by('tarih', 'sube__ad', 'urun_anahtar', 'olusturma'))
    return {'subeler': list(subeler), 'ozet': ozet_liste, 'genel': genel, 'urunler': urunler,
            'gunler': gunler, 'detay': detay}


def sube_adi(k):
    return k.sube.ad if k.sube_id else (k.sube_ad_arsiv or '—')


def giren_adi(k):
    return k.personel_ad_arsiv or (k.personel.ad_soyad if k.personel_id else '') or '—'


# ---- Excel ------------------------------------------------------------------------
def excel_bytes(veri, etiket, sube_etiketi):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    wb = openpyxl.Workbook()
    mavi = PatternFill('solid', fgColor='162AA3')
    zebra = PatternFill('solid', fgColor='F4F6FB')
    ince = Side(style='thin', color='D7DCEB')
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
    bas_font = Font(bold=True, color='FFFFFF')

    def baslik_satiri(ws, satir, basliklar):
        for i, b in enumerate(basliklar, 1):
            c = ws.cell(row=satir, column=i, value=b)
            c.fill, c.font, c.border = mavi, bas_font, kenar
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def genislik(ws, olculer):
        for i, w in enumerate(olculer, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # 1) Özet — şubeler arası karşılaştırma
    ws = wb.active
    ws.title = 'Şube Özeti'
    ws['A1'] = 'Mutfak Zayi Raporu'
    ws['A1'].font = Font(bold=True, size=14, color='162AA3')
    ws['A2'] = 'Dönem: %s   ·   %s' % (etiket, sube_etiketi)
    baslik_satiri(ws, 4, ['Şube', 'Kayıt sayısı', 'Ağırlık (kg)', 'Hacim (lt)', 'Adet'])
    r = 5
    for o in veri['ozet']:
        vals = [o['sube'].ad, o['kayit'], float(o['top']['agirlik'] / 1000),
                float(o['top']['hacim'] / 1000), float(o['top']['adet'])]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = kenar
            if (r % 2) == 0:
                c.fill = zebra
            if i >= 3:
                c.number_format = '#,##0.00' if i < 5 else '#,##0'
        r += 1
    g = veri['genel']
    for i, v in enumerate(['TOPLAM', g['kayit'], float(g['top']['agirlik'] / 1000),
                           float(g['top']['hacim'] / 1000), float(g['top']['adet'])], 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font, c.border = Font(bold=True), kenar
        if i >= 3:
            c.number_format = '#,##0.00' if i < 5 else '#,##0'
    ws['A%d' % (r + 2)] = 'Ağırlık gr+kg, hacim ml+lt toplamıdır; farklı ölçüler birbirine eklenmez.'
    ws['A%d' % (r + 2)].font = Font(italic=True, color='5B6472')
    genislik(ws, [34, 14, 16, 16, 12])
    ws.freeze_panes = 'A5'
    excel_logo(ws)

    # 2) Ürün x Şube (sayısal; gr / ml / adet cinsinden)
    ws2 = wb.create_sheet('Ürün x Şube')
    ws2['A1'] = 'Ürün bazında zayi — %s' % etiket
    ws2['A1'].font = Font(bold=True, size=13, color='162AA3')
    ws2['A2'] = 'Miktarlar taban birimdedir: ağırlık = gr, hacim = ml, adet = adet.'
    ws2['A2'].font = Font(italic=True, color='5B6472')
    basliklar = ['Ürün', 'Ölçü', 'Birim'] + [s.ad for s in veri['subeler']] + ['TOPLAM']
    baslik_satiri(ws2, 4, basliklar)
    taban_birim = {'agirlik': 'gr', 'hacim': 'ml', 'adet': 'adet'}
    r = 5
    for u in veri['urunler']:
        satir = [u['ad'], OLCU_ETIKET[u['olcu']], taban_birim[u['olcu']]]
        satir += [float(u['subeler'].get(s.id, 0)) if u['subeler'].get(s.id) else None for s in veri['subeler']]
        satir.append(float(u['toplam']))
        for i, v in enumerate(satir, 1):
            c = ws2.cell(row=r, column=i, value=v)
            c.border = kenar
            if (r % 2) == 0:
                c.fill = zebra
            if i > 3:
                c.number_format = '#,##0.##'
        ws2.cell(row=r, column=len(satir)).font = Font(bold=True)
        r += 1
    genislik(ws2, [30, 12, 8] + [16] * len(veri['subeler']) + [14])
    ws2.freeze_panes = 'D5'
    excel_logo(ws2)

    # 3) Gün gün
    ws3 = wb.create_sheet('Gün Gün')
    baslik_satiri(ws3, 1, ['Tarih', 'Kayıt sayısı', 'Ağırlık (kg)', 'Hacim (lt)', 'Adet'])
    r = 2
    for gn in veri['gunler']:
        vals = [gn['tarih'], gn['kayit'], float(gn['top']['agirlik'] / 1000),
                float(gn['top']['hacim'] / 1000), float(gn['top']['adet'])]
        for i, v in enumerate(vals, 1):
            c = ws3.cell(row=r, column=i, value=v)
            c.border = kenar
            if i == 1:
                c.number_format = 'DD.MM.YYYY'
            elif i >= 3:
                c.number_format = '#,##0.00' if i < 5 else '#,##0'
        r += 1
    genislik(ws3, [14, 14, 16, 16, 12])
    ws3.freeze_panes = 'A2'
    excel_logo(ws3)

    # 4) Detay — her kayıt
    ws4 = wb.create_sheet('Kayıtlar')
    baslik_satiri(ws4, 1, ['Tarih', 'Şube', 'Ürün', 'Miktar', 'Birim', 'Giren', 'Açıklama'])
    r = 2
    for k in veri['detay']:
        vals = [k.tarih, sube_adi(k), k.urun, float(k.miktar), k.birim, giren_adi(k), k.aciklama]
        for i, v in enumerate(vals, 1):
            c = ws4.cell(row=r, column=i, value=v)
            c.border = kenar
            c.alignment = Alignment(vertical='top', wrap_text=(i == 7))
            if i == 1:
                c.number_format = 'DD.MM.YYYY'
            elif i == 4:
                c.number_format = '#,##0.##'
        r += 1
    genislik(ws4, [12, 26, 28, 12, 8, 24, 50])
    ws4.freeze_panes = 'A2'
    ws4.auto_filter.ref = 'A1:G%d' % max(r - 1, 1)
    excel_logo(ws4)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- PDF --------------------------------------------------------------------------
def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from django.conf import settings
    adaylar = [os.path.join(settings.BASE_DIR, 'static', 'fonts'),
               '/usr/share/fonts/truetype/dejavu']
    for base in adaylar:
        try:
            pdfmetrics.registerFont(TTFont('DejaVu', os.path.join(base, 'DejaVuSans.ttf')))
            pdfmetrics.registerFont(TTFont('DejaVu-Bold', os.path.join(base, 'DejaVuSans-Bold.ttf')))
            return 'DejaVu', 'DejaVu-Bold'
        except Exception:
            continue
    return 'Helvetica', 'Helvetica-Bold'


def pdf_bytes(veri, etiket, sube_etiketi):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf

    FONT, FONTB = _pdf_font()
    BRAND = colors.HexColor('#162AA3')
    BRAND_050 = colors.HexColor('#EEF1FB')
    ZEBRA = colors.HexColor('#F4F6FB')
    LINE = colors.HexColor('#E3E6EF')
    MUTED = colors.HexColor('#6B7280')

    st_baslik = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_alt = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_sb = ParagraphStyle('sb', fontName=FONTB, fontSize=11.5, textColor=BRAND, leading=15)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=8.2, leading=10.4)
    st_cb = ParagraphStyle('cb', fontName=FONTB, fontSize=8.2, leading=10.4)
    st_cw = ParagraphStyle('cw', fontName=FONTB, fontSize=8.2, leading=10.4, textColor=colors.white)
    st_not = ParagraphStyle('n', fontName=FONT, fontSize=7.6, textColor=MUTED, leading=10)

    LEFT = RIGHT = 15 * mm
    W = A4[0] - LEFT - RIGHT

    def P(t, st=st_c):
        return Paragraph(escape(str(t)), st)

    def tablo(rows, kolonlar, son_toplam=False, hizala_sag=()):
        t = Table(rows, colWidths=kolonlar, repeatRows=1)
        stil = [
            ('BACKGROUND', (0, 0), (-1, 0), BRAND),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2 if son_toplam else -1), [colors.white, ZEBRA]),
            ('GRID', (0, 0), (-1, -1), 0.3, LINE),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ]
        if son_toplam:
            stil += [('BACKGROUND', (0, -1), (-1, -1), BRAND_050), ('LINEABOVE', (0, -1), (-1, -1), 0.7, BRAND)]
        for c in hizala_sag:
            stil.append(('ALIGN', (c, 0), (c, -1), 'RIGHT'))
        t.setStyle(TableStyle(stil))
        return t

    el = [Paragraph('Mutfak Zayi Raporu', st_baslik),
          Paragraph('Dönem: <b>%s</b> · %s' % (escape(etiket), escape(sube_etiketi)), st_alt),
          Spacer(1, 8)]

    # Şube özeti
    el.append(Paragraph('Şubeler arası karşılaştırma', st_sb))
    el.append(HRFlowable(width='100%', thickness=1, color=LINE, spaceBefore=2, spaceAfter=4))
    rows = [[P(h, st_cw) for h in ('Şube', 'Kayıt', 'Ağırlık', 'Hacim', 'Adet')]]
    for o in veri['ozet']:
        rows.append([P(o['sube'].ad), P(o['kayit']), P(o['yazi']['agirlik'] or '—'),
                     P(o['yazi']['hacim'] or '—'), P(o['yazi']['adet'] or '—')])
    g = veri['genel']
    rows.append([P('TOPLAM', st_cb), P(g['kayit'], st_cb), P(g['yazi']['agirlik'] or '—', st_cb),
                 P(g['yazi']['hacim'] or '—', st_cb), P(g['yazi']['adet'] or '—', st_cb)])
    el.append(tablo(rows, [W * .36, W * .10, W * .18, W * .18, W * .18], son_toplam=True, hizala_sag=(1, 2, 3, 4)))
    el.append(Spacer(1, 3))
    el.append(Paragraph('Ağırlık gr+kg, hacim ml+lt toplamıdır; farklı ölçüler birbirine eklenmez.', st_not))
    el.append(Spacer(1, 10))

    # Şube şube: ürün toplamları
    for s in veri['subeler']:
        sube_urunleri = [u for u in veri['urunler'] if u['subeler'].get(s.id)]
        el.append(Paragraph(s.ad, st_sb))
        el.append(HRFlowable(width='100%', thickness=1, color=LINE, spaceBefore=2, spaceAfter=4))
        if not sube_urunleri:
            el.append(Paragraph('Bu dönemde zayi kaydı yok.', st_alt))
            el.append(Spacer(1, 8))
            continue
        rows = [[P(h, st_cw) for h in ('Ürün', 'Toplam zayi')]]
        for u in sube_urunleri:
            rows.append([P(u['ad']), P(miktar_yaz(u['subeler'][s.id], u['olcu']))])
        el.append(tablo(rows, [W * .65, W * .35], hizala_sag=(1,)))
        el.append(Spacer(1, 10))

    # Kayıt dökümü
    if veri['detay']:
        el.append(Paragraph('Kayıt dökümü', st_sb))
        el.append(HRFlowable(width='100%', thickness=1, color=LINE, spaceBefore=2, spaceAfter=4))
        rows = [[P(h, st_cw) for h in ('Tarih', 'Şube', 'Ürün', 'Miktar', 'Giren', 'Açıklama')]]
        for k in veri['detay']:
            rows.append([P(k.tarih.strftime('%d.%m.%Y')), P(sube_adi(k)), P(k.urun), P(kalem_miktar_yaz(k)),
                         P(giren_adi(k)), P(k.aciklama or '')])
        el.append(tablo(rows, [W * .12, W * .17, W * .20, W * .11, W * .16, W * .24], hizala_sag=(3,)))

    el.append(Spacer(1, 8))
    el.append(Paragraph('Bu rapor otomatik olarak oluşturulmuştur · geekpanel.net', st_not))
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=A4, left_margin=LEFT, right_margin=RIGHT, font=FONT,
              title='Mutfak Zayi Raporu')
    return buf.getvalue()
