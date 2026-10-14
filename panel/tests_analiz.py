"""Sevkiyat talep analizi:  python manage.py test panel.tests_analiz"""
import io
from decimal import Decimal as D

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from . import sevkiyat_analiz as sa
from .models import Personel, Rol, SevkiyatKalem, SevkiyatTalep, Sube, Urun

_AYAR = dict(
    SECURE_SSL_REDIRECT=False,
    STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
              'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}},
)


def kisi(rol, sube=None):
    u = User.objects.create(username='u%d' % (User.objects.count() + 1))
    return Personel.objects.create(user=u, ad_soyad=rol, rol=rol, sube=sube)


@override_settings(**_AYAR)
class AnalizTestleri(TestCase):
    def setUp(self):
        self.a = Sube.objects.create(ad='A Şube')
        self.b = Sube.objects.create(ad='B Şube')
        self.depo = Sube.objects.create(ad='Depo', depo_mu=True)
        self.sut = Urun.objects.create(ad='Süt', kategori='Süt', birim='LİTRE', ust_birim='KOLİ', koli_icerigi=12)
        self.ay = timezone.localdate().replace(day=1)

    def talep(self, sube, durum, kalemler):
        t = SevkiyatTalep.objects.create(sube=sube, durum=durum, olusturan_ad='x')
        for ist, birim, sev in kalemler:
            SevkiyatKalem.objects.create(talep=t, urun=self.sut, urun_ad='Süt', koli_icerigi=12,
                                         istenen_miktar=D(ist), istenen_birim=birim,
                                         sevkiyat_miktar=None if sev is None else D(sev), sevkiyat_birim=birim)
        return t

    def veri(self, subeler=None):
        return sa.analiz_verisi(self.ay, subeler or [self.a, self.b])

    def test_koli_tabana_cevrilir_ve_toplanir(self):
        self.talep(self.a, 'Onaylandı', [(2, 'KOLİ', 2)])       # 24 litre
        self.talep(self.a, 'Onaylandı', [(10, 'LİTRE', 8)])     # 10 litre
        u = self.veri()['urunler'][0]
        self.assertEqual((u['birim'], u['talep'], u['giden']), ('LİTRE', D(34), D(32)))

    def test_giden_yalniz_onayli_bekleyen_ayri(self):
        self.talep(self.a, 'Onaylandı', [(5, 'LİTRE', 4)])
        self.talep(self.a, 'Talep', [(7, 'LİTRE', None)])
        u = self.veri()['urunler'][0]
        self.assertEqual((u['talep'], u['giden'], u['bekleyen']), (D(12), D(4), D(7)))
        self.assertEqual(u['karsilama'], 80)

    def test_sube_kirilimi(self):
        self.talep(self.a, 'Onaylandı', [(5, 'LİTRE', 5)])
        self.talep(self.b, 'Onaylandı', [(9, 'LİTRE', 3)])
        u = self.veri()['urunler'][0]
        self.assertEqual(u['subeler'][self.a.id]['talep'], D(5))
        self.assertEqual(u['subeler'][self.b.id]['giden'], D(3))

    def test_kapsam_disi_sube_ve_diger_ay_sayilmaz(self):
        self.talep(self.b, 'Onaylandı', [(5, 'LİTRE', 5)])
        self.assertEqual(self.veri([self.a])['urunler'], [])
        t = self.talep(self.a, 'Onaylandı', [(5, 'LİTRE', 5)])
        SevkiyatTalep.objects.filter(pk=t.pk).update(olusturma=timezone.now() - timezone.timedelta(days=70))
        self.assertEqual(self.veri([self.a])['urunler'], [])

    def test_roller_ve_kapsam(self):
        self.assertEqual({s.id for s in sa.analiz_subeleri(kisi(Rol.GENEL_MUDUR))}, {self.a.id, self.b.id})  # depo yok
        mudur = kisi(Rol.MUDUR); mudur.sorumlu_subeler.add(self.a)
        self.assertEqual([s.id for s in sa.analiz_subeleri(mudur)], [self.a.id])
        self.assertEqual(sa.analiz_subeleri(kisi(Rol.SEF, self.a)), [])

    def test_sayfa_yetkisi(self):
        self.talep(self.a, 'Onaylandı', [(5, 'LİTRE', 5)])
        for rol, kod in ((Rol.GENEL_MUDUR, 200), (Rol.OPERATOR, 200), (Rol.SATIN_ALMA, 200), (Rol.YATIRIMCI, 200),
                         (Rol.SEF, 302), (Rol.PERSONEL, 302), (Rol.SEVKIYAT, 302)):
            self.client.force_login(kisi(rol, self.a).user)
            self.assertEqual(self.client.get(reverse('sevkiyat_analiz')).status_code, kod, rol)
            self.assertEqual(self.client.get(reverse('sevkiyat_analiz_rapor') + '?bicim=pdf').status_code,
                             200 if kod == 200 else 302, rol)

    def test_mudur_baska_subeyi_goremez(self):
        self.talep(self.b, 'Onaylandı', [(5, 'LİTRE', 5)])
        m = kisi(Rol.MUDUR); m.sorumlu_subeler.add(self.a)
        self.client.force_login(m.user)
        r = self.client.get(reverse('sevkiyat_analiz') + '?sube=%d' % self.b.id)   # yetkisiz şube -> kendi kapsamı
        self.assertNotContains(r, 'B Şube')
        self.assertContains(r, 'Bu ayda')

    def test_excel_ve_pdf(self):
        self.talep(self.a, 'Onaylandı', [(5, 'LİTRE', 5)])
        self.talep(self.b, 'Talep', [(3, 'KOLİ', None)])
        self.client.force_login(kisi(Rol.GENEL_MUDUR).user)
        x = self.client.get(reverse('sevkiyat_analiz_rapor') + '?bicim=xlsx')
        self.assertEqual(x.status_code, 200)
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ['Ürün Özeti', 'Şube x Ürün', 'Şube Özeti'])
        satirlar = [[c.value for c in r] for r in wb['Ürün Özeti'].iter_rows(min_row=6)]
        self.assertEqual(satirlar[0][:5], ['Süt', 'LİTRE', 41.0, 5.0, 36.0])
        p = self.client.get(reverse('sevkiyat_analiz_rapor') + '?bicim=pdf&sube=%d' % self.a.id)
        self.assertEqual(p['Content-Type'], 'application/pdf')
        self.assertTrue(p.content.startswith(b'%PDF'))

    def test_bozuk_ay_ve_bos_veri_hata_vermez(self):
        self.client.force_login(kisi(Rol.GENEL_MUDUR).user)
        for q in ('?ay=abc', '?ay=2026-13', '?ay=1999-01', '?sube=zzz&urun=x'):
            self.assertEqual(self.client.get(reverse('sevkiyat_analiz') + q).status_code, 200, q)
        self.assertEqual(self.client.get(reverse('sevkiyat_analiz_rapor') + '?bicim=pdf').status_code, 200)
        self.assertEqual(self.client.get(reverse('sevkiyat_analiz_rapor') + '?bicim=xlsx').status_code, 200)
