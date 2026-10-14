"""Sevkiyat analizi — ay bazında şube / ürün: ne kadar talep edildi, ne kadar gitti.

Tanımlar
  Talep    : şubenin sipariş ettiği miktar (istenen_miktar), tüm siparişler.
  Giden    : çıkışı ONAYLANMIŞ (Onaylandı / Teslim Edildi) siparişlerin sevkiyat miktarı.
  Bekleyen : henüz çıkışı onaylanmamış siparişlerin talep miktarı (satın almada, sevkiyatta, onay bekliyor...).

Birimler: koli/paket/set gibi "üst birim" ile girilen miktarlar ürünün taban birimine çevrilir
(1 koli = N adet/litre...), böylece aynı ürün farklı birimlerle girilse de toplanır. Farklı ürünler
birbirine eklenmez; her ürün kendi biriminde gösterilir.
"""
import calendar
import datetime
import io
from decimal import Decimal, ROUND_HALF_UP

from django.utils import timezone

from .models import Rol, SevkiyatDurumu, SevkiyatKalem, Sube

# Analizi görebilen roller (Bölge Müdürü yalnızca sorumlu olduğu şubeleri görür)
ANALIZ_ROLLERI = (Rol.GENEL_MUDUR, Rol.OPERATOR, Rol.YATIRIMCI, Rol.SATIN_ALMA, Rol.MUDUR)

_GIDEN_DURUMLAR = (SevkiyatDurumu.ONAYLANDI, SevkiyatDurumu.TESLIM)
_SIFIR = Decimal('0')
_AYLAR = ['Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos', 'Eylül', 'Ekim', 'Kasım', 'Aralık']


def ay_oku(metin):
    """'2026-10' -> date(2026,10,1); bozuksa bu ay."""
    try:
        y, m = (int(x) for x in (metin or '').split('-'))
        return datetime.date(y, m, 1) if 1 <= m <= 12 and 2000 <= y <= 2100 else timezone.localdate().replace(day=1)
    except (ValueError, TypeError):
        return timezone.localdate().replace(day=1)


def ay_etiketi(ay_ilk):
    return '%s %d' % (_AYLAR[ay_ilk.month - 1], ay_ilk.year)


def ay_komsulari(ay_ilk):
    onceki = (ay_ilk - datetime.timedelta(days=1)).replace(day=1)
    sonraki = (ay_ilk + datetime.timedelta(days=31)).replace(day=1)
    return onceki, sonraki


def sayi_yaz(d):
    """Decimal -> '1.250,5' (Türkçe biçim, gereksiz sıfır yok)."""
    if d is None:
        return '0'
    d = Decimal(d).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    s = '{:,.2f}'.format(d).replace(',', '#').replace('.', ',').replace('#', '.')
    if ',' in s:
        s = s.rstrip('0').rstrip(',')
    return s


def analiz_subeleri(personel):
    """Bu kişinin analizde görebileceği şubeler (depolar hariç)."""
    if personel is None or personel.rol not in ANALIZ_ROLLERI:
        return []
    if personel.rol == Rol.MUDUR:
        qs = personel.sorumlu_subeler.filter(depo_mu=False)
    else:
        qs = Sube.objects.filter(depo_mu=False)
    return list(qs.order_by('ad'))


def _taban(k, miktar, birim):
    u = k.urun
    if u is not None and u.ust_birim and u.ust_birim != u.birim and birim == u.ust_birim:
        return miktar * Decimal(k.koli_icerigi or u.koli_icerigi or 1), u.birim
    return miktar, birim


def _yeni():
    return {'talep': _SIFIR, 'giden': _SIFIR, 'bekleyen': _SIFIR, 'siparis': set()}


def analiz_verisi(ay_ilk, subeler):
    """subeler: görülebilecek Sube listesi (zaten yetkiye göre süzülmüş)."""
    sube_ids = [s.id for s in subeler]
    son = ay_ilk.replace(day=calendar.monthrange(ay_ilk.year, ay_ilk.month)[1]) + datetime.timedelta(days=1)
    tz = timezone.get_current_timezone()
    bas = timezone.make_aware(datetime.datetime.combine(ay_ilk, datetime.time.min), tz)
    bit = timezone.make_aware(datetime.datetime.combine(son, datetime.time.min), tz)

    kalemler = (SevkiyatKalem.objects
                .filter(talep__sube_id__in=sube_ids, talep__olusturma__gte=bas, talep__olusturma__lt=bit)
                .select_related('talep', 'urun'))

    urunler = {}
    sube_ozet = {s.id: {'sube': s, 'siparis': set(), 'onayli': set(), 'bekleyen': set(), 'kalem': 0}
                 for s in subeler}
    for k in kalemler:
        t = k.talep
        so = sube_ozet.get(t.sube_id)
        if so is None:
            continue
        onayli = t.durum in _GIDEN_DURUMLAR
        so['siparis'].add(t.id)
        (so['onayli'] if onayli else so['bekleyen']).add(t.id)
        so['kalem'] += 1

        talep_m, birim = _taban(k, k.istenen_miktar or _SIFIR, k.istenen_birim)
        giden_m = _SIFIR
        if onayli:
            if k.sevkiyat_miktar is not None:
                g, gb = k.sevkiyat_miktar, (k.sevkiyat_birim or k.istenen_birim)
            elif k.satinalma_miktar is not None:
                g, gb = k.satinalma_miktar, (k.satinalma_birim or k.istenen_birim)
            else:
                g, gb = k.istenen_miktar or _SIFIR, k.istenen_birim
            giden_m, _gb = _taban(k, g, gb)
        ad = (k.urun_ad or (k.urun.ad if k.urun_id else '') or '—').strip()
        anahtar = '%s|%s' % (ad.casefold(), birim)
        u = urunler.get(anahtar)
        if u is None:
            ust = k.urun.ust_birim if (k.urun_id and k.urun.ust_birim and k.urun.ust_birim != k.urun.birim) else ''
            icerik = (k.koli_icerigi or (k.urun.koli_icerigi if k.urun_id else 1) or 1) if ust else 1
            u = urunler[anahtar] = {'anahtar': anahtar, 'ad': ad, 'birim': birim, 'ust': ust, 'icerik': icerik,
                                    'kategori': k.kategori or '', **_yeni(), 'subeler': {}}
        for hedef in (u, u['subeler'].setdefault(t.sube_id, _yeni())):
            hedef['talep'] += talep_m
            hedef['siparis'].add(t.id)
            if onayli:
                hedef['giden'] += giden_m
            else:
                hedef['bekleyen'] += talep_m

    liste = sorted(urunler.values(), key=lambda x: (-x['talep'], x['ad'].casefold()))
    for u in liste:
        u['siparis_adet'] = len(u['siparis'])
        u['karsilama'] = _yuzde(u['giden'], u['talep'] - u['bekleyen'])
        if u['ust'] and u['icerik'] > 1:
            u['talep_ust'] = u['talep'] / Decimal(u['icerik'])
            u['giden_ust'] = u['giden'] / Decimal(u['icerik'])
        else:
            u['talep_ust'] = u['giden_ust'] = None
        for sv in u['subeler'].values():
            sv['karsilama'] = _yuzde(sv['giden'], sv['talep'] - sv['bekleyen'])

    ozet = []
    for s in subeler:
        so = sube_ozet[s.id]
        cesit = sum(1 for u in liste if s.id in u['subeler'])
        ozet.append({'sube': s, 'siparis': len(so['siparis']), 'onayli': len(so['onayli']),
                     'bekleyen': len(so['bekleyen']), 'kalem': so['kalem'], 'cesit': cesit})
    return {'ay': ay_ilk, 'etiket': ay_etiketi(ay_ilk), 'urunler': liste, 'ozet': ozet,
            'subeler': list(subeler),
            'toplam_siparis': sum(o['siparis'] for o in ozet),
            'toplam_onayli': sum(o['onayli'] for o in ozet),
            'toplam_bekleyen': sum(o['bekleyen'] for o in ozet)}


def _yuzde(giden, hedef):
    """Onaylanmış siparişlerin talebine göre karşılanma yüzdesi (hedef 0 ise None)."""
    if not hedef or hedef <= 0:
        return None
    return int((giden / hedef * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def gorunum(veri, sube=None, urun_anahtar=None, ilk=10):
    """Sayfa için: seçili şubeye göre ürün listesi + grafik verisi.

    sube: Sube ya da None (tüm şubeler toplamı).
    Dönen: satirlar (tablo), grafik_urun (ilk N ürün), grafik_sube (seçili ürünün şube kırılımı), secili_urun.
    """
    satirlar = []
    for u in veri['urunler']:
        if sube is None:
            v = u
        else:
            v = u['subeler'].get(sube.id)
            if not v:
                continue
        satirlar.append({'anahtar': u['anahtar'], 'ad': u['ad'], 'birim': u['birim'], 'ust': u['ust'],
                         'icerik': u['icerik'], 'talep': v['talep'], 'giden': v['giden'],
                         'bekleyen': v['bekleyen'], 'karsilama': v['karsilama'],
                         'siparis': len(v['siparis'])})
    satirlar.sort(key=lambda x: (-x['talep'], x['ad'].casefold()))

    def yazilar(x):
        x['t_yaz'], x['g_yaz'], x['b_yaz'] = sayi_yaz(x['talep']), sayi_yaz(x['giden']), sayi_yaz(x['bekleyen'])
        return x

    for x in satirlar:
        yazilar(x)

    def grafik(liste, mx_alani):
        mx = max([max(g['talep'], g['giden']) for g in liste] or [_SIFIR]) or Decimal('1')
        for g in liste:
            g['t_yuzde'] = float(g['talep'] / mx * 100)
            g['g_yuzde'] = float(g['giden'] / mx * 100)
        return liste

    grafik_urun = grafik([dict(x) for x in satirlar[:ilk]], None)

    secili = next((u for u in veri['urunler'] if u['anahtar'] == urun_anahtar), None)
    if secili is None and veri['urunler']:
        secili = veri['urunler'][0]
    grafik_sube = []
    if secili is not None:
        for s in veri['subeler']:
            v = secili['subeler'].get(s.id) or {'talep': _SIFIR, 'giden': _SIFIR, 'bekleyen': _SIFIR}
            grafik_sube.append({'ad': s.ad, 'talep': v['talep'], 'giden': v['giden'], 'bekleyen': v['bekleyen']})
        grafik_sube = [g for g in grafik_sube if g['talep'] or g['giden']]
        grafik_sube.sort(key=lambda g: -g['talep'])
        grafik(grafik_sube, None)
        for x in grafik_sube:
            yazilar(x)
    return {'satirlar': satirlar, 'grafik_urun': grafik_urun, 'grafik_sube': grafik_sube,
            'secili_urun': secili}


# ----------------------------------------------------------------------------- Excel
def excel_bytes(veri, sube=None):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    wb = openpyxl.Workbook()
    mavi = PatternFill('solid', fgColor='162AA3')
    zebra = PatternFill('solid', fgColor='F4F6FB')
    ince = Side(style='thin', color='D7DCEB')
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)

    def sayfa(ws, baslik, basliklar, satirlar, olculer, sayisal=()):
        excel_logo(ws)
        ws['C2'] = baslik
        ws['C2'].font = Font(bold=True, size=14, color='162AA3')
        ws['C3'] = '%s · %s' % (veri['etiket'], sube.ad if sube else 'Tüm şubeler')
        ws['C3'].font = Font(color='6B7280')
        ust = 5
        for i, b in enumerate(basliklar, 1):
            c = ws.cell(row=ust, column=i, value=b)
            c.fill, c.font, c.border = mavi, Font(bold=True, color='FFFFFF'), kenar
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for r, satir in enumerate(satirlar, ust + 1):
            for i, v in enumerate(satir, 1):
                c = ws.cell(row=r, column=i, value=v)
                c.border = kenar
                if (r - ust) % 2 == 0:
                    c.fill = zebra
                if i in sayisal:
                    c.number_format = '#,##0.##'
                    c.alignment = Alignment(horizontal='right')
        for i, w in enumerate(olculer, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=ust + 1, column=1)

    def f(d):
        return float(d)

    # 1) Ürün özeti (seçili şubeye göre)
    g = gorunum(veri, sube, ilk=0)
    ws = wb.active
    ws.title = 'Ürün Özeti'
    sayfa(ws, 'Sevkiyat Analizi — Ürün Özeti',
          ['Ürün', 'Birim', 'Talep edilen', 'Giden (onaylı)', 'Bekleyen', 'Karşılama %', 'Sipariş sayısı'],
          [[s['ad'], s['birim'], f(s['talep']), f(s['giden']), f(s['bekleyen']),
            s['karsilama'] if s['karsilama'] is not None else '', s['siparis']] for s in g['satirlar']],
          [34, 10, 15, 15, 13, 13, 13], sayisal=(3, 4, 5))

    # 2) Şube x ürün
    ws = wb.create_sheet('Şube x Ürün')
    satirlar = []
    for s in veri['subeler']:
        if sube is not None and s.id != sube.id:
            continue
        for u in veri['urunler']:
            v = u['subeler'].get(s.id)
            if not v:
                continue
            satirlar.append([s.ad, u['ad'], u['birim'], f(v['talep']), f(v['giden']), f(v['bekleyen']),
                             v['karsilama'] if v['karsilama'] is not None else '', len(v['siparis'])])
    sayfa(ws, 'Sevkiyat Analizi — Şube x Ürün',
          ['Şube', 'Ürün', 'Birim', 'Talep edilen', 'Giden (onaylı)', 'Bekleyen', 'Karşılama %', 'Sipariş sayısı'],
          satirlar, [24, 34, 10, 15, 15, 13, 13, 13], sayisal=(4, 5, 6))

    # 3) Şube özeti
    ws = wb.create_sheet('Şube Özeti')
    sayfa(ws, 'Sevkiyat Analizi — Şube Özeti',
          ['Şube', 'Sipariş', 'Onaylı (gitti)', 'Bekleyen', 'Kalem', 'Ürün çeşidi'],
          [[o['sube'].ad, o['siparis'], o['onayli'], o['bekleyen'], o['kalem'], o['cesit']]
           for o in veri['ozet'] if sube is None or o['sube'].id == sube.id],
          [26, 12, 15, 12, 10, 14])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- PDF
def pdf_bytes(veri, sube=None):
    from xml.sax.saxutils import escape
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf
    from .zayi_rapor import _pdf_font

    FONT, FONTB = _pdf_font()
    BRAND = colors.HexColor('#162AA3')
    CREMA = colors.HexColor('#F2B544')
    ZEBRA = colors.HexColor('#F4F6FB')
    LINE = colors.HexColor('#E3E6EF')
    MUTED = colors.HexColor('#6B7280')
    BRAND_050 = colors.HexColor('#EEF1FB')

    st_baslik = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_alt = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_sb = ParagraphStyle('sb', fontName=FONTB, fontSize=11.5, textColor=BRAND, leading=15)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=8.2, leading=10.4)
    st_cw = ParagraphStyle('cw', fontName=FONTB, fontSize=8.2, leading=10.4, textColor=colors.white)
    st_not = ParagraphStyle('n', fontName=FONT, fontSize=7.6, textColor=MUTED, leading=10)

    LEFT = RIGHT = 15 * mm
    W = A4[0] - LEFT - RIGHT

    def P(t, st=st_c):
        return Paragraph(escape(str(t)), st)

    def tablo(rows, kolonlar, sag=()):
        t = Table(rows, colWidths=kolonlar, repeatRows=1)
        stil = [('BACKGROUND', (0, 0), (-1, 0), BRAND),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, ZEBRA]),
                ('GRID', (0, 0), (-1, -1), 0.3, LINE), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4)]
        stil += [('ALIGN', (c, 0), (c, -1), 'RIGHT') for c in sag]
        t.setStyle(TableStyle(stil))
        return t

    def grafik(satirlar, baslik):
        """İlk 8 ürün — talep (mavi) / giden (krem) sütunları."""
        satirlar = satirlar[:8]
        d = Drawing(W, 150)
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 30, 38, W - 40, 92
        bc.data = [[float(s['talep']) for s in satirlar], [float(s['giden']) for s in satirlar]]
        bc.categoryAxis.categoryNames = [(s['ad'][:12] + '…') if len(s['ad']) > 12 else s['ad'] for s in satirlar]
        bc.categoryAxis.labels.fontName = FONT
        bc.categoryAxis.labels.fontSize = 6.5
        bc.categoryAxis.labels.angle = 25
        bc.categoryAxis.labels.dy = -8
        bc.categoryAxis.labels.boxAnchor = 'ne'
        bc.valueAxis.labels.fontName = FONT
        bc.valueAxis.labels.fontSize = 6.5
        bc.valueAxis.valueMin = 0
        bc.valueAxis.gridStrokeColor = LINE
        bc.valueAxis.visibleGrid = True
        bc.bars[0].fillColor, bc.bars[1].fillColor = BRAND, CREMA
        bc.bars[0].strokeColor = bc.bars[1].strokeColor = None
        bc.barSpacing = 1
        bc.groupSpacing = 8
        d.add(bc)
        d.add(String(30, 140, baslik, fontName=FONTB, fontSize=8, fillColor=MUTED))
        d.add(String(W - 120, 140, '■ Talep', fontName=FONT, fontSize=7.5, fillColor=BRAND))
        d.add(String(W - 70, 140, '■ Giden', fontName=FONT, fontSize=7.5, fillColor=CREMA))
        return d

    def deger(x, ust=None):
        return sayi_yaz(x)

    el = [Paragraph('Sevkiyat Analizi', st_baslik),
          Paragraph('Dönem: <b>%s</b> · %s' % (escape(veri['etiket']), escape(sube.ad if sube else 'Tüm şubeler')), st_alt),
          Paragraph('Sipariş: %d · Onaylı (gitti): %d · Bekleyen: %d' % (
              sum(o['siparis'] for o in veri['ozet'] if sube is None or o['sube'].id == sube.id),
              sum(o['onayli'] for o in veri['ozet'] if sube is None or o['sube'].id == sube.id),
              sum(o['bekleyen'] for o in veri['ozet'] if sube is None or o['sube'].id == sube.id)), st_alt),
          Spacer(1, 8)]

    bolumler = [(sube.ad if sube else 'Tüm şubeler', sube)]
    if sube is None:
        bolumler += [(s.ad, s) for s in veri['subeler']]
    for ad, sb in bolumler:
        g = gorunum(veri, sb, ilk=8)
        baslik = [Paragraph(ad, st_sb), HRFlowable(width='100%', thickness=1, color=LINE, spaceBefore=2, spaceAfter=4)]
        if not g['satirlar']:
            el.append(KeepTogether(baslik + [Paragraph('Bu ayda sipariş yok.', st_alt), Spacer(1, 8)]))
            continue
        el.append(KeepTogether(baslik + [grafik(g['satirlar'], 'En çok talep edilen ürünler (her ürün kendi biriminde)')]))
        rows = [[P(h, st_cw) for h in ('Ürün', 'Birim', 'Talep', 'Giden', 'Bekleyen', 'Karşılama')]]
        for s in g['satirlar']:
            rows.append([P(s['ad']), P(s['birim']), P(deger(s['talep'])), P(deger(s['giden'])),
                         P(deger(s['bekleyen'])), P('%%%d' % s['karsilama'] if s['karsilama'] is not None else '—')])
        el.append(tablo(rows, [W * .34, W * .12, W * .14, W * .14, W * .14, W * .12], sag=(2, 3, 4, 5)))
        el.append(Spacer(1, 12))

    el.append(Paragraph('Talep: tüm siparişler. Giden: çıkışı onaylanmış siparişler. Bekleyen: henüz onaylanmamış siparişlerin '
                        'talebi. Koli/paket gibi üst birimler ürünün taban birimine çevrilmiştir. Karşılama = giden ÷ '
                        '(onaylı siparişlerin talebi).', st_not))
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=A4, left_margin=LEFT, right_margin=RIGHT, font=FONT, title='Sevkiyat Analizi')
    return buf.getvalue()
