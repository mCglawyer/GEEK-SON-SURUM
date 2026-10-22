"""Tüketim & maliyet raporlarının Excel / PDF çıktıları.

İki rapor türü:
  ay    : Aylık kullanım ve maliyet (açılış + gelen − kapanış). Sayım ay sonunda yapıldığı için kullanım ay bazındadır.
  gelen : Seçilen tarih aralığında (gün / hafta / ay / özel) şubelere ulaşan mal ve tutarı. Gün gün dökülür.
"""
import datetime
import io
from decimal import Decimal
from xml.sax.saxutils import escape

from django.utils import timezone

from . import tuketim as tk
from .models import SevkiyatBirim

_SIFIR = Decimal('0')


def aralik_oku(bicim, ref, bas=None, bit=None):
    """bicim: gun | hafta | ay | ozel  -> (bas, bit, etiket). ref: referans gün."""
    ref = ref or timezone.localdate()
    if bicim == 'gun':
        return ref, ref, ref.strftime('%d.%m.%Y')
    if bicim == 'hafta':
        b = ref - datetime.timedelta(days=ref.weekday())
        e = b + datetime.timedelta(days=6)
        return b, e, '%s – %s' % (b.strftime('%d.%m.%Y'), e.strftime('%d.%m.%Y'))
    if bicim == 'ozel' and bas and bit:
        if bit < bas:
            bas, bit = bit, bas
        bit = min(bit, bas + datetime.timedelta(days=366))
        return bas, bit, '%s – %s' % (bas.strftime('%d.%m.%Y'), bit.strftime('%d.%m.%Y'))
    b = ref.replace(day=1)
    e = (b + datetime.timedelta(days=32)).replace(day=1) - datetime.timedelta(days=1)
    return b, e, tk.ay_etiketi(b)


def gelen_hesapla(subeler, bas, bit, fiyatli=True):
    """Aralıkta gelen mal. Satır: gun, sube, urun, miktar(taban), birim, fiyat, tutar. fiyatsız ürün tutarsız kalır."""
    ids = [s.id for s in subeler]
    sube_by = {s.id: s for s in subeler}
    defter = tk.FiyatDefteri() if fiyatli else None
    satirlar = []
    if not ids:
        return {'satirlar': [], 'toplam_tutar': _SIFIR, 'fiyat_eksik': []}
    eksik = set()
    for k in tk._sevk_kalemleri(ids, bas - datetime.timedelta(days=1), bit):
        gun = timezone.localtime(k._eff).date()
        if not (bas <= gun <= bit):
            continue
        m, b = tk._gelen_miktar(k)
        taban = tk.taban_miktar(k, m, b)
        if not taban:
            continue
        sube = sube_by[k.talep.sube_id]
        ad = (k.urun_ad or (k.urun.ad if k.urun_id else '—')).strip()
        birim = k.urun.birim if k.urun_id else (b or '')
        fiyat = defter.bul(k.urun_id, sube, gun) if (defter and k.urun_id) else None
        tutar = tk._yuvarla(taban * fiyat) if fiyat is not None else None
        if fiyatli and fiyat is None:
            eksik.add(ad)
        satirlar.append({'gun': gun, 'sube': sube, 'urun': ad, 'miktar': taban, 'birim': str(birim or ''),
                         'fiyat': fiyat, 'tutar': tutar})
    satirlar.sort(key=lambda r: (r['gun'], r['sube'].ad, r['urun']))
    return {'satirlar': satirlar, 'fiyat_eksik': sorted(eksik),
            'toplam_tutar': sum((r['tutar'] for r in satirlar if r['tutar'] is not None), _SIFIR)}


def _sayi(v, n=2):
    if v is None:
        return '—'
    s = '{:,.{n}f}'.format(Decimal(v), n=n).rstrip('0').rstrip('.') if n else '{:,.0f}'.format(Decimal(v))
    return s.replace(',', 'X').replace('.', ',').replace('X', '.')


def _para(v):
    return '—' if v is None else '{:,.2f}'.format(Decimal(v)).replace(',', 'X').replace('.', ',').replace('X', '.') + ' ₺'


def _baslik(tur, etiket, secili):
    ad = 'Aylık tüketim ve maliyet' if tur == 'ay' else 'Gelen mal ve tutar'
    return ad, '%s · %s' % (etiket, secili.ad if secili else 'Tüm şubeler')


# ----------------------------------------------------------------------------- Excel
def excel_bytes(tur, etiket, secili, veri=None, gelen=None):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    wb = openpyxl.Workbook()
    mavi = PatternFill('solid', fgColor='162AA3')
    zebra = PatternFill('solid', fgColor='F4F6FB')
    ince = Side(style='thin', color='D7DCEB')
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
    baslik, alt = _baslik(tur, etiket, secili)

    def sayfa(ws, basliklar, satirlar, olculer, para=(), toplam=None, not_metni=''):
        excel_logo(ws)
        ws['C2'] = baslik
        ws['C2'].font = Font(bold=True, size=14, color='162AA3')
        ws['C3'] = alt
        ws['C3'].font = Font(color='6B7280')
        ust = 5
        for i, b in enumerate(basliklar, 1):
            c = ws.cell(row=ust, column=i, value=b)
            c.fill, c.font, c.border = mavi, Font(bold=True, color='FFFFFF'), kenar
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for r, satir in enumerate(satirlar, ust + 1):
            for i, v in enumerate(satir, 1):
                c = ws.cell(row=r, column=i, value=(float(v) if isinstance(v, Decimal) else v))
                c.border = kenar
                if (r - ust) % 2 == 0:
                    c.fill = zebra
                if isinstance(v, (Decimal, float, int)) and not isinstance(v, bool):
                    c.number_format = '#,##0.00' if i in para else '#,##0.##'
                    c.alignment = Alignment(horizontal='right')
        son = ust + len(satirlar)
        if toplam:
            for i, v in enumerate(toplam, 1):
                c = ws.cell(row=son + 1, column=i, value=(float(v) if isinstance(v, Decimal) else v))
                c.font, c.border = Font(bold=True), kenar
                if isinstance(v, (Decimal, float)):
                    c.number_format = '#,##0.00'
            son += 1
        for i in range(1, len(basliklar) + 1):
            ws.column_dimensions[get_column_letter(i)].width = olculer[i - 1] if i <= len(olculer) else 12
        ws.freeze_panes = ws.cell(row=ust + 1, column=1)
        if not_metni:
            ws.cell(row=son + 2, column=1, value=not_metni).font = Font(italic=True, color='6B7280')

    if tur == 'ay':
        ws = wb.active
        ws.title = 'Şube özeti'
        sayfa(ws, ['Şube', 'Durum', 'Tutar (TL)'],
              [[c['sube'].ad, {'tam': 'Hesaplandı', 'sayim_bekleniyor': 'Ay sonu sayımı bekleniyor',
                               'acilis_yok': 'Önceki ay sayımı yok', 'baslamadi': 'Sistem başlamadı'}[c['durum']],
                c.get('tutar')] for c in veri['subeler']],
              [30, 30, 16], para=(3,), toplam=['TOPLAM', '', veri['genel_tutar']],
              not_metni='Kullanım = Açılış (önceki ay sonu sayımı) + Gelen (sevkiyat) − Kapanış (bu ay sonu sayımı). '
                        'Kullanıma fire de dahildir.')
        ws = wb.create_sheet('Ürün toplamı')
        sayfa(ws, ['Ürün', 'Açılış', 'Gelen', 'Kapanış', 'Kullanım', 'Birim maliyet (TL)', 'Tutar (TL)'],
              [[r['urun'].ad, r['ilk'], r['gelen'], r['son'], r['kullanim'], r.get('birim_maliyet'), r.get('tutar')]
               for r in veri['toplam']], [40, 12, 12, 12, 12, 16, 16], para=(6, 7),
              toplam=['TOPLAM', '', '', '', '', '', veri['genel_tutar']],
              not_metni='Miktarlar sayım biriminde (ml, gr, adet).')
        ws = wb.create_sheet('Şube × ürün')
        satirlar = []
        for c in veri['subeler']:
            for r in c['satirlar']:
                satirlar.append([c['sube'].ad, r['urun'].ad, r['ilk'], r['gelen'], r['son'], r['kullanim'],
                                 r.get('birim_maliyet'), r.get('tutar')])
        sayfa(ws, ['Şube', 'Ürün', 'Açılış', 'Gelen', 'Kapanış', 'Kullanım', 'Birim maliyet (TL)', 'Tutar (TL)'],
              satirlar, [26, 40, 12, 12, 12, 12, 16, 16], para=(7, 8))
    else:
        ws = wb.active
        ws.title = 'Gün gün'
        sayfa(ws, ['Tarih', 'Şube', 'Ürün', 'Miktar', 'Birim', 'Birim fiyat (TL)', 'Tutar (TL)'],
              [[r['gun'].strftime('%d.%m.%Y'), r['sube'].ad, r['urun'], r['miktar'], r['birim'], r['fiyat'], r['tutar']]
               for r in gelen['satirlar']], [12, 26, 40, 12, 10, 16, 16], para=(6, 7),
              toplam=['TOPLAM', '', '', '', '', '', gelen['toplam_tutar']],
              not_metni='Çıkışı onaylanan sevkiyatlar; mal teslim gününün fiyatıyla değerlenir. Miktar taban birimdedir (adet/kg).')
        ws = wb.create_sheet('Şube toplamı')
        toplam = {}
        for r in gelen['satirlar']:
            t = toplam.setdefault(r['sube'].id, [r['sube'].ad, _SIFIR, 0])
            t[1] += r['tutar'] or _SIFIR
            t[2] += 1
        sayfa(ws, ['Şube', 'Tutar (TL)', 'Kalem sayısı'], sorted(toplam.values()), [30, 16, 14], para=(2,),
              toplam=['TOPLAM', gelen['toplam_tutar'], ''])
        ws = wb.create_sheet('Ürün toplamı')
        ut = {}
        for r in gelen['satirlar']:
            t = ut.setdefault((r['urun'], r['birim']), [r['urun'], r['birim'], _SIFIR, _SIFIR])
            t[2] += r['miktar']
            t[3] += r['tutar'] or _SIFIR
        sayfa(ws, ['Ürün', 'Birim', 'Miktar', 'Tutar (TL)'], sorted(ut.values()), [40, 10, 14, 16], para=(4,))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- PDF
def pdf_bytes(tur, etiket, secili, veri=None, gelen=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf
    from .zayi_rapor import _pdf_font

    FONT, FONTB = _pdf_font()
    BRAND, ZEBRA = colors.HexColor('#162AA3'), colors.HexColor('#F4F6FB')
    LINE, MUTED = colors.HexColor('#E3E6EF'), colors.HexColor('#6B7280')
    sayfa_boyu = landscape(A4)
    LEFT = RIGHT = 12 * mm
    W = sayfa_boyu[0] - LEFT - RIGHT
    st_b = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_a = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_sb = ParagraphStyle('sb', fontName=FONTB, fontSize=11, textColor=BRAND, leading=15)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=8, leading=10.4)
    st_cw = ParagraphStyle('cw', fontName=FONTB, fontSize=8, leading=10.4, textColor=colors.white)
    st_n = ParagraphStyle('n', fontName=FONT, fontSize=7.4, textColor=MUTED, leading=10)

    def P(t, st=st_c):
        return Paragraph(escape('' if t is None else str(t)), st)

    def tablo(rows, genislik, sag_ilk=1, son_kalin=False):
        t = Table(rows, colWidths=genislik, repeatRows=1)
        stil = [('BACKGROUND', (0, 0), (-1, 0), BRAND), ('GRID', (0, 0), (-1, -1), .4, LINE),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ALIGN', (sag_ilk, 1), (-1, -1), 'RIGHT'),
                ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
        for i in range(2, len(rows), 2):
            stil.append(('BACKGROUND', (0, i), (-1, i), ZEBRA))
        t.setStyle(TableStyle(stil))
        return t

    baslik, alt = _baslik(tur, etiket, secili)
    el = [Paragraph(escape(baslik), st_b), Paragraph(escape(alt), st_a), Spacer(1, 8)]
    if tur == 'ay':
        rows = [[P('Şube', st_cw), P('Durum', st_cw), P('Tutar', st_cw)]]
        ad = {'tam': 'Hesaplandı', 'sayim_bekleniyor': 'Ay sonu sayımı bekleniyor', 'acilis_yok': 'Önceki ay sayımı yok',
              'baslamadi': 'Sistem başlamadı'}
        for c in veri['subeler']:
            rows.append([P(c['sube'].ad), P(ad[c['durum']]), P(_para(c.get('tutar')))])
        rows.append([P('TOPLAM'), P(''), P(_para(veri['genel_tutar']))])
        el += [Paragraph('Şubelere göre', st_sb), tablo(rows, [W * .5, W * .3, W * .2], 2), Spacer(1, 8)]
        rows = [[P(h, st_cw) for h in ('Ürün', 'Açılış', 'Gelen', 'Kapanış', 'Kullanım', 'Birim maliyet', 'Tutar')]]
        for r in veri['toplam']:
            rows.append([P(r['urun'].ad), P(_sayi(r['ilk'])), P(_sayi(r['gelen'])), P(_sayi(r['son'])),
                         P(_sayi(r['kullanim'])), P(_para(r.get('birim_maliyet'))), P(_para(r.get('tutar')))])
        el += [Paragraph('Ürün bazında (sayım biriminde)', st_sb),
               tablo(rows, [W * .30, W * .10, W * .10, W * .10, W * .12, W * .14, W * .14]), Spacer(1, 6)]
        if veri.get('fiyat_eksik'):
            el.append(Paragraph('Fiyatı girilmemiş ürünler toplama dahil değil: ' + escape(', '.join(veri['fiyat_eksik'])), st_n))
        el.append(Paragraph('Kullanım = Açılış + Gelen − Kapanış. Fire kullanımın içindedir. Maliyet ağırlıklı ortalama fiyatla bulunur.', st_n))
    else:
        rows = [[P(h, st_cw) for h in ('Tarih', 'Şube', 'Ürün', 'Miktar', 'Birim fiyat', 'Tutar')]]
        for r in gelen['satirlar']:
            rows.append([P(r['gun'].strftime('%d.%m.%Y')), P(r['sube'].ad), P(r['urun']),
                         P('%s %s' % (_sayi(r['miktar']), r['birim'].lower())), P(_para(r['fiyat'])), P(_para(r['tutar']))])
        rows.append([P('TOPLAM'), P(''), P(''), P(''), P(''), P(_para(gelen['toplam_tutar']))])
        if len(rows) == 2:
            rows.insert(1, [P('Bu aralıkta gelen mal yok.')] + [P('')] * 5)
        el.append(tablo(rows, [W * .10, W * .20, W * .30, W * .12, W * .13, W * .15], 3))
        if gelen['fiyat_eksik']:
            el += [Spacer(1, 6), Paragraph('Fiyatı girilmemiş ürünler toplama dahil değil: ' + escape(', '.join(gelen['fiyat_eksik'])), st_n)]
        el.append(Paragraph('Çıkışı onaylanan sevkiyatlar, teslim gününün fiyatıyla değerlenir.', st_n))
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=sayfa_boyu, left_margin=LEFT, right_margin=RIGHT, font=FONT, title=baslik)
    return buf.getvalue()
