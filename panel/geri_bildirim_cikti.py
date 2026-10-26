"""Geri Bildirim Yönetimi için Excel / PDF çıktıları."""
import io
from xml.sax.saxutils import escape


def excel_bytes(bildirimler, etiket):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from .excel_logo import excel_logo

    mavi = PatternFill('solid', fgColor='162AA3'); kenar = Border(*[Side(style='thin', color='D7DCEB')] * 4)
    bas_font = Font(bold=True, color='FFFFFF')
    wb = Workbook(); ws = wb.active; ws.title = 'Geri bildirimler'
    ws['A1'] = 'Geri bildirimler'; ws['A1'].font = Font(bold=True, size=14, color='162AA3')
    ws['A2'] = etiket; ws['A2'].font = Font(italic=True, color='5B6472')
    kolonlar = ['Tarih', 'Kategori', 'Durum', 'Şube', 'Metin', 'Yönetim notu']
    for j, k in enumerate(kolonlar, 1):
        c = ws.cell(row=4, column=j, value=k); c.fill, c.font, c.border = mavi, bas_font, kenar
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, b in enumerate(bildirimler, 5):
        satir = [b.olusturma.strftime('%d.%m.%Y %H:%M'), b.kategori, b.durum, b.sube.ad if b.sube else '—', b.metin, b.yonetim_notu]
        for j, v in enumerate(satir, 1):
            c = ws.cell(row=i, column=j, value=v); c.border = kenar
            c.alignment = Alignment(vertical='top', wrap_text=True)
    genislik = [16, 12, 14, 20, 60, 40]
    for j in range(1, len(kolonlar) + 1):
        ws.column_dimensions[get_column_letter(j)].width = genislik[j - 1]
    ws.freeze_panes = 'A5'
    excel_logo(ws)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def pdf_bytes(bildirimler, etiket):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from .pdf_letterhead import build_pdf
    from .zayi_rapor import _pdf_font

    FONT, FONTB = _pdf_font()
    BRAND = colors.HexColor('#162AA3'); MUTED = colors.HexColor('#5b6472'); LINE = colors.HexColor('#d7dceb'); SOFT = colors.HexColor('#EEF1FB')
    st_b = ParagraphStyle('b', fontName=FONTB, fontSize=14, textColor=BRAND, leading=18)
    st_a = ParagraphStyle('a', fontName=FONT, fontSize=9.5, textColor=MUTED, leading=14)
    st_c = ParagraphStyle('c', fontName=FONT, fontSize=8, leading=11)
    P = lambda t, s=st_c: Paragraph(escape(str(t)), s)

    data = [[P(k, ParagraphStyle('h', fontName=FONTB, fontSize=8, leading=10, textColor=colors.white))
              for k in ['Tarih', 'Kategori', 'Durum', 'Şube', 'Metin']]]
    for b in bildirimler:
        data.append([P(b.olusturma.strftime('%d.%m.%Y %H:%M')), P(b.kategori), P(b.durum), P(b.sube.ad if b.sube else '—'), P(b.metin)])
    t = Table(data, colWidths=[24 * mm, 20 * mm, 20 * mm, 30 * mm, 88 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), BRAND), ('GRID', (0, 0), (-1, -1), 0.4, LINE), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, SOFT]), ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    el = [P('Geri bildirimler', st_b), P(etiket, st_a), Spacer(1, 4 * mm), t]
    buf = io.BytesIO()
    build_pdf(buf, el, pagesize=A4, left_margin=12 * mm, right_margin=12 * mm, font=FONT, title='Geri bildirimler')
    return buf.getvalue()
