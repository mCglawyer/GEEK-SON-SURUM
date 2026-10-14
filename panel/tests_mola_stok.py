"""Mola (QR okutunca başlar, 15/45 sonradan) ve stok ay seçimi:  python manage.py test panel.tests_mola_stok"""
import datetime

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .models import (Bildirim, MolaOturum, MolaQRAyar, Personel, Rol, StokSayim, StokUrun, Sube, SubeMolaToken)

_AYAR = dict(
    SECURE_SSL_REDIRECT=False,
    STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
              'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}},
)


def kisi(rol, sube=None, ad=None):
    u = User.objects.create(username='u%d' % (User.objects.count() + 1))
    return Personel.objects.create(user=u, ad_soyad=ad or rol, rol=rol, sube=sube)


@override_settings(**_AYAR)
class MolaTestleri(TestCase):
    """QR okutunca başlar, tekrar okutunca biter; günlük toplam."""

    def setUp(self):
        self.sube = Sube.objects.create(ad='Merkez')
        MolaQRAyar.objects.create(acik=True)
        SubeMolaToken.objects.create(sube=self.sube, token='tok123')
        self.p = kisi(Rol.PERSONEL, self.sube, 'Ayşe')
        self.client.force_login(self.p.user)
        self.url = reverse('mola_qr_giris', args=['tok123'])

    def okut(self, **kw):
        return self.client.get(self.url, follow=True, **kw)

    def acik(self):
        return MolaOturum.objects.filter(personel=self.p, bitis__isnull=True)

    def geri_al(self, dakika, saniye=0):
        """Açık molanın başlangıcını geçmişe çek (zaman geçmiş gibi)."""
        MolaOturum.objects.filter(personel=self.p, bitis__isnull=True).update(
            baslangic=timezone.now() - datetime.timedelta(minutes=dakika, seconds=saniye))

    def test_qr_okutunca_mola_hemen_baslar(self):
        r = self.okut()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.acik().count(), 1)
        self.assertContains(r, 'Moladasın')
        self.assertContains(r, 'başladı')
        self.assertEqual(self.acik().get().sure_dk, 60)          # günlük hak
        self.assertNotContains(r, 'Molayı Bitir')                # buton yok

    def test_baslayinca_teyit_bildirimi_gider(self):
        self.okut()
        b = Bildirim.objects.filter(alici=self.p).get()
        self.assertIn('başladı', b.mesaj)
        self.assertIn('60 dk', b.mesaj)

    def test_ilk_60_saniyede_ikinci_okutma_bitirmez(self):
        self.okut()
        r = self.okut()
        self.assertEqual(self.acik().count(), 1)
        self.assertContains(r, 'zaten')
        self.assertEqual(MolaOturum.objects.filter(personel=self.p).count(), 1)

    def test_donunce_okutunca_direkt_biter(self):
        self.okut(); self.geri_al(23)
        r = self.okut()
        self.assertEqual(self.acik().count(), 0)
        m = MolaOturum.objects.get(personel=self.p)
        self.assertEqual(m.kullanilan_dk, 23)
        self.assertContains(r, 'Molan bitti')
        self.assertContains(r, '23 dk')
        self.assertTrue(Bildirim.objects.filter(alici=self.p, mesaj__contains='bitti').exists())

    def test_bitiriken_yeni_mola_acilmaz_ve_durum_sayfasi_yenilenebilir(self):
        self.okut(); self.geri_al(10); self.okut()
        for _ in range(3):
            self.client.get(reverse('mola_durum'))
        self.assertEqual(MolaOturum.objects.filter(personel=self.p).count(), 1)
        self.assertEqual(self.acik().count(), 0)

    def test_bitirdikten_hemen_sonra_okutma_yeni_mola_acmaz(self):
        self.okut(); self.geri_al(10); self.okut()
        r = self.okut()                                           # çift okutma
        self.assertEqual(MolaOturum.objects.filter(personel=self.p).count(), 1)
        self.assertContains(r, 'az önce')

    def test_gun_icinde_birden_cok_mola_ve_toplam(self):
        self.okut(); self.geri_al(15); self.okut()                # 15 dk
        MolaOturum.objects.filter(personel=self.p).update(
            bitis=timezone.now() - datetime.timedelta(minutes=5), baslangic=timezone.now() - datetime.timedelta(minutes=20))
        self.okut(); self.geri_al(30)
        r = self.okut()                                           # 30 dk
        self.assertEqual(MolaOturum.objects.filter(personel=self.p).count(), 2)
        self.assertContains(r, '45 dk')                           # günlük toplam
        self.assertContains(r, 'Bugünkü molalarım')

    def test_kalan_hak_dusulur(self):
        self.okut(); self.geri_al(40); self.okut()
        MolaOturum.objects.filter(personel=self.p).update(bitis=timezone.now() - datetime.timedelta(minutes=3))
        self.okut()
        self.assertEqual(self.acik().get().sure_dk, 20)           # 60 - 40

    def test_unutulan_mola_otomatik_kapanir_ve_yenisi_baslar(self):
        self.okut(); self.geri_al(200)
        r = self.okut()
        eski = MolaOturum.objects.filter(personel=self.p, bitis__isnull=False).get()
        self.assertEqual(eski.kullanilan_dk, 120)
        self.assertIn('Otomatik', eski.manuel_not)
        self.assertEqual(self.acik().count(), 1)
        self.assertContains(r, 'otomatik kapatıldı')

    def test_arka_plan_gorevi_unutulani_kapatir(self):
        from . import mola_kurallari as mk
        self.okut(); self.geri_al(130)
        self.assertEqual(mk.eskileri_kapat(), 1)
        self.assertEqual(self.acik().count(), 0)
        self.geri_al(1)
        self.assertEqual(mk.eskileri_kapat(), 0)

    def test_qr_kapaliyken_mola_acilmaz(self):
        MolaQRAyar.objects.update(acik=False)
        self.okut()
        self.assertEqual(MolaOturum.objects.count(), 0)

    def test_gecersiz_token_mola_acmaz(self):
        self.client.get(reverse('mola_qr_giris', args=['yok']))
        self.assertEqual(MolaOturum.objects.count(), 0)

    def test_baska_siteden_gelen_istek_mola_baslatmaz(self):
        self.client.get(self.url, HTTP_SEC_FETCH_SITE='cross-site')
        self.assertEqual(MolaOturum.objects.count(), 0)

    def test_eski_butonlu_sayfa_sure_sec_degistirmez(self):
        self.okut()
        self.client.post(self.url, {'islem': 'sure_sec', 'sure': '15'})
        self.client.post(self.url, {'islem': 'baslat', 'sure': '15'})
        m = self.acik().get()
        self.assertEqual(m.sure_dk, 60)
        self.assertEqual(MolaOturum.objects.filter(personel=self.p).count(), 1)

    def test_eski_sayfadan_bitir_calisir(self):
        self.okut(); self.geri_al(5)
        self.client.post(self.url, {'islem': 'bitir'})
        self.assertEqual(self.acik().count(), 0)

    def test_baska_subenin_qr_ile_de_bitirilebilir(self):
        diger = Sube.objects.create(ad='Diğer')
        SubeMolaToken.objects.create(sube=diger, token='baska')
        self.okut(); self.geri_al(5)
        self.client.get(reverse('mola_qr_giris', args=['baska']))
        self.assertEqual(self.acik().count(), 0)
        self.assertEqual(MolaOturum.objects.get(personel=self.p).sube, self.sube)

    def test_mutfak_sorumlusu_mola_kullanmaz(self):
        ms = kisi(Rol.MUTFAK_SORUMLUSU, self.sube, 'Sor')
        self.client.force_login(ms.user)
        self.client.get(self.url)
        self.assertEqual(MolaOturum.objects.filter(personel=ms).count(), 0)

    def test_eszamanli_iki_okutma_cift_mola_acmaz(self):
        from . import mola_kurallari as mk
        simdi = timezone.now()
        mk.okut(self.p, self.sube, simdi)
        mk.okut(self.p, self.sube, simdi + datetime.timedelta(seconds=2))
        self.assertEqual(MolaOturum.objects.filter(personel=self.p).count(), 1)
        self.assertEqual(self.acik().count(), 1)


@override_settings(**_AYAR)
class MolaRaporTestleri(TestCase):
    def setUp(self):
        self.sube = Sube.objects.create(ad='Merkez')
        self.p = kisi(Rol.PERSONEL, self.sube, 'Ayşe')
        self.gm = kisi(Rol.GENEL_MUDUR, None, 'GM')
        bugun = timezone.localtime().replace(hour=12, minute=0, second=0, microsecond=0)
        for bas, dk in ((0, 20), (180, 50)):                      # 12:00 (20 dk) ve 15:00 (50 dk) = 70 dk
            b = bugun + datetime.timedelta(minutes=bas)
            MolaOturum.objects.create(personel=self.p, sube=self.sube, sure_dk=60, baslangic=b,
                                      bitis=b + datetime.timedelta(minutes=dk), kullanilan_dk=dk)
        self.client.force_login(self.gm.user)

    def test_gunluk_toplam_ve_saatler(self):
        from . import mola_kurallari as mk
        g = mk.gunluk_gruplar(MolaOturum.objects.select_related('personel', 'sube'))
        self.assertEqual(len(g), 1)
        self.assertEqual(g[0]['toplam_dk'], 70)
        self.assertTrue(g[0]['asti'])
        self.assertEqual([m['giris'].strftime('%H:%M') for m in g[0]['molalar']], ['12:00', '15:00'])

    def test_gecmis_sayfasi_gun_bazinda_gosterir(self):
        r = self.client.get(reverse('mola_gecmis') + '?hizli=bugun')
        self.assertContains(r, '70 dk')
        self.assertContains(r, '12:00–12:20')
        self.assertContains(r, '15:00–15:50')
        self.assertContains(r, 'Aşım')

    def test_excel_ve_pdf(self):
        import io
        import openpyxl
        x = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&bicim=xlsx')
        self.assertEqual(x.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ['Günlük Özet', 'Tüm Molalar'])
        satir = [c.value for c in next(wb['Günlük Özet'].iter_rows(min_row=6, max_row=6))]
        self.assertEqual((satir[1], satir[3], satir[4], satir[6]), ('Ayşe', 2, 70, 'Aşım'))
        self.assertIn('12:00–12:20', satir[7])
        p = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&bicim=pdf')
        self.assertTrue(p.content.startswith(b'%PDF'))

    def _ek_kisi(self, ad, sube=None):
        k = kisi(Rol.PERSONEL, sube or self.sube, ad)
        b = timezone.localtime().replace(hour=9, minute=0, second=0, microsecond=0)
        MolaOturum.objects.create(personel=k, sube=k.sube, sure_dk=60, baslangic=b,
                                  bitis=b + datetime.timedelta(minutes=10), kullanilan_dk=10)
        return k

    def test_arama_turkce_harf_duyarsiz(self):
        self._ek_kisi('İbrahim Şahin')
        r = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&q=ibrahim')
        self.assertContains(r, 'İbrahim Şahin')
        self.assertNotContains(r, 'Ayşe</b>')
        r = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&q=SAHIN')
        self.assertContains(r, 'İbrahim Şahin')

    def test_durum_asim_suzgeci_ve_sayaclar(self):
        self._ek_kisi('Kısa Mola')
        r = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&durum=asim')
        self.assertContains(r, 'Ayşe')
        self.assertNotContains(r, 'Kısa Mola')
        self.assertEqual(r.context['toplam_grup'], 1)
        self.assertEqual(r.context['toplam_kayit'], 2)

    def test_rapor_suzgece_uyar(self):
        import io
        import openpyxl
        self._ek_kisi('Kısa Mola')
        x = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&durum=asim&bicim=xlsx')
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        adlar = [r[1].value for r in wb['Günlük Özet'].iter_rows(min_row=6) if r[1].value]
        self.assertEqual(adlar, ['Ayşe'])

    def test_sayfalama(self):
        for i in range(25):
            self._ek_kisi('Kişi %02d' % i)
        r = self.client.get(reverse('mola_gecmis') + '?hizli=bugun')
        self.assertEqual(r.context['toplam_grup'], 26)
        self.assertEqual(len(r.context['gruplar']), 20)
        self.assertEqual(r.context['sayfa_sayisi'], 2)
        r2 = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&sayfa=2')
        self.assertEqual(len(r2.context['gruplar']), 6)
        r3 = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&sayfa=99')   # sınır dışı -> son sayfa
        self.assertEqual(r3.context['sayfa_no'], 2)
        self.assertEqual(self.client.get(reverse('mola_gecmis') + '?sayfa=abc&durum=zzz&sirala=x').status_code, 200)

    def test_siralama_en_uzun(self):
        self._ek_kisi('Kısa Mola')
        r = self.client.get(reverse('mola_gecmis') + '?hizli=bugun&sirala=uzun')
        self.assertEqual(r.context['gruplar'][0]['ad'], 'Ayşe')

    def test_personel_raporu_goremez(self):
        self.client.force_login(self.p.user)
        self.assertEqual(self.client.get(reverse('mola_gecmis')).status_code, 302)

    def test_izleme_json_unutulani_kapatir(self):
        MolaOturum.objects.create(personel=self.p, sube=self.sube, sure_dk=60,
                                  baslangic=timezone.now() - datetime.timedelta(hours=5))
        r = self.client.get(reverse('mola_izleme_json'))
        self.assertEqual(r.json()['molalar'], [])
        self.assertEqual(MolaOturum.objects.filter(bitis__isnull=True).count(), 0)


@override_settings(**_AYAR)
class StokAyTestleri(TestCase):
    def setUp(self):
        self.sube = Sube.objects.create(ad='Merkez')
        self.sef = kisi(Rol.SEF, self.sube)
        self.urun = StokUrun.objects.create(ad='Süt', kapali_icerik=12)
        self.client.force_login(self.sef.user)
        bu = timezone.localdate().replace(day=1)
        self.bu = bu
        self.onceki = (bu - datetime.timedelta(days=1)).replace(day=1)

    def kaydet(self, kayit_ay=None):
        d = {'islem': 'stok_kaydet', 'kapali_%d' % self.urun.id: '3'}
        if kayit_ay:
            d['kayit_ay'] = kayit_ay
        return self.client.post(reverse('stok'), d)

    def test_onay_penceresi_iki_ayi_sunar(self):
        r = self.client.get(reverse('stok'))
        self.assertContains(r, 'Bu hangi ayın stoğu?')
        self.assertContains(r, 'data-ay="%s"' % self.bu.strftime('%Y-%m'))
        self.assertContains(r, 'data-ay="%s"' % self.onceki.strftime('%Y-%m'))

    def test_onceki_ay_secilirse_oraya_yazilir(self):
        self.kaydet(self.onceki.strftime('%Y-%m'))
        self.assertTrue(StokSayim.objects.filter(sube=self.sube, ay=self.onceki).exists())
        self.assertFalse(StokSayim.objects.filter(sube=self.sube, ay=self.bu).exists())

    def test_bu_ay_secilirse_bu_aya_yazilir(self):
        self.kaydet(self.bu.strftime('%Y-%m'))
        self.assertTrue(StokSayim.objects.filter(sube=self.sube, ay=self.bu).exists())

    def test_izinsiz_ay_secimi_yok_sayilir(self):
        self.kaydet('2019-01')
        self.assertFalse(StokSayim.objects.filter(ay=datetime.date(2019, 1, 1)).exists())
        self.assertTrue(StokSayim.objects.filter(sube=self.sube, ay=self.bu).exists())

    def test_secim_yoksa_ekrandaki_aya_yazar(self):
        self.kaydet()
        self.assertTrue(StokSayim.objects.filter(sube=self.sube, ay=self.bu).exists())

    def test_var_olan_kayit_isaretlenir(self):
        StokSayim.objects.create(sube=self.sube, ay=self.onceki)
        r = self.client.get(reverse('stok'))
        self.assertContains(r, 'kayıt var, güncellenir')
